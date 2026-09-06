"""
Admin Dashboard Service
=======================

Contains all business logic used by the Admin Dashboard.

The service is intentionally separate from the API router so that
the same functionality can later be reused by other interfaces.
"""

from fastapi import HTTPException, status

import base64
import binascii
import csv
import io
import re
from sqlalchemy import func, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from app.models.user import User
from app.models.role import Role
from app.models.admin_profile import AdminProfile
from app.services.auth_service import hash_password


# ============================================================
# HELPER: NORMALIZE ACCOUNT STATUS
# ============================================================

def _normalize_status(value) -> str:
    """
    Convert account status values/enums into plain uppercase strings.
    """

    if value is None:
        return ""

    value = str(value)

    if "." in value:
        value = value.split(".")[-1]

    return value.upper()


# ============================================================
# HELPER: USER LIST RESPONSE
# ============================================================

def _user_to_list_item(user: User) -> dict:
    """
    Convert a User SQLAlchemy object into the structure expected
    by AdminUserListItem.
    """

    role_name = None

    if user.role is not None:
        role_name = str(user.role.role_name)

        if "." in role_name:
            role_name = role_name.split(".")[-1]

        role_name = role_name.upper()

    return {
        "user_id": user.user_id,
        "email": user.email,
        "full_name": user.full_name,
        "phone": user.phone,
        "role_name": role_name or "",
        "account_status": _normalize_status(user.account_status),
        "created_at": user.created_at,
        "updated_at": user.updated_at,
        "last_login_at": user.last_login_at,
    }


# ============================================================
# HELPER: PROFILE RESPONSE
# ============================================================

def _user_to_detail(user: User) -> dict:
    """
    Convert a User SQLAlchemy object into the structure expected
    by AdminUserDetailResponse.
    """

    role_name = None

    if user.role is not None:
        role_name = str(user.role.role_name)

        if "." in role_name:
            role_name = role_name.split(".")[-1]

        role_name = role_name.upper()

    student_profile = None

    if user.student_profile is not None:
        profile = user.student_profile

        student_profile = {
            "student_id": profile.student_id,
            "enrollment_number": profile.enrollment_number,
            "roll_number": profile.roll_number,
            "date_of_birth": profile.date_of_birth,
        }

    faculty_profile = None

    if user.faculty_profile is not None:
        profile = user.faculty_profile

        faculty_profile = {
            "faculty_id": profile.faculty_id,
            "employee_id": profile.employee_id,
            "department": profile.department,
            "designation": profile.designation,
        }

    admin_profile = None

    if user.admin_profile is not None:
        profile = user.admin_profile

        admin_profile = {
            "admin_id": profile.admin_id,
            "employee_id": profile.employee_id,
            "designation": profile.designation,
        }

    return {
        "user_id": user.user_id,

        "email": user.email,
        "full_name": user.full_name,
        "phone": user.phone,

        "role_name": role_name or "",

        "account_status": _normalize_status(
            user.account_status
        ),

        "created_at": user.created_at,
        "updated_at": user.updated_at,

        "last_login_at": user.last_login_at,

        "student_profile": student_profile,
        "faculty_profile": faculty_profile,
        "admin_profile": admin_profile,
    }


# ============================================================
# DASHBOARD STATISTICS
# ============================================================

def get_dashboard_statistics(
    db: Session,
) -> dict:
    """
    Return summary statistics for the Admin Dashboard.
    """

    # --------------------------------------------------------
    # TOTAL USERS
    # --------------------------------------------------------

    total_users = db.scalar(
        select(func.count(User.user_id))
    ) or 0

    # --------------------------------------------------------
    # STATUS COUNTS
    # --------------------------------------------------------

    pending_users = db.scalar(
        select(func.count(User.user_id))
        .where(User.account_status == "PENDING")
    ) or 0

    active_users = db.scalar(
        select(func.count(User.user_id))
        .where(User.account_status == "ACTIVE")
    ) or 0

    suspended_users = db.scalar(
        select(func.count(User.user_id))
        .where(User.account_status == "SUSPENDED")
    ) or 0

    inactive_users = db.scalar(
        select(func.count(User.user_id))
        .where(User.account_status == "INACTIVE")
    ) or 0

    # --------------------------------------------------------
    # ROLE COUNTS
    # --------------------------------------------------------

    total_students = db.scalar(
        select(func.count(User.user_id))
        .join(Role)
        .where(Role.role_name == "STUDENT")
    ) or 0

    total_faculty = db.scalar(
        select(func.count(User.user_id))
        .join(Role)
        .where(Role.role_name == "FACULTY")
    ) or 0

    total_admins = db.scalar(
        select(func.count(User.user_id))
        .join(Role)
        .where(Role.role_name == "ADMIN")
    ) or 0

    return {
        "total_users": total_users,

        "pending_users": pending_users,
        "active_users": active_users,
        "suspended_users": suspended_users,
        "inactive_users": inactive_users,

        "total_students": total_students,
        "total_faculty": total_faculty,
        "total_admins": total_admins,
    }


# ============================================================
# LIST USERS
# ============================================================

def list_users(
    db: Session,
    role_name: str | None = None,
    account_status: str | None = None,
):
    """
    Return users for the Admin Dashboard.

    Optional filters:

        role_name
        account_status
    """

    statement = (
        select(User)
        .options(
            joinedload(User.role)
        )
        .join(User.role)
    )

    # --------------------------------------------------------
    # ROLE FILTER
    # --------------------------------------------------------

    if role_name:

        normalized_role = role_name.strip().upper()

        allowed_roles = {
            "ADMIN",
            "FACULTY",
            "STUDENT",
        }

        if normalized_role not in allowed_roles:

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Invalid role filter. "
                    "Allowed values: ADMIN, FACULTY, STUDENT."
                ),
            )

        statement = statement.where(
            Role.role_name == normalized_role
        )

    # --------------------------------------------------------
    # STATUS FILTER
    # --------------------------------------------------------

    if account_status:

        normalized_status = (
            account_status.strip().upper()
        )

        allowed_statuses = {
            "PENDING",
            "ACTIVE",
            "SUSPENDED",
            "INACTIVE",
        }

        if normalized_status not in allowed_statuses:

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Invalid account status filter. "
                    "Allowed values: "
                    "PENDING, ACTIVE, SUSPENDED, INACTIVE."
                ),
            )

        statement = statement.where(
            User.account_status == normalized_status
        )

    statement = statement.order_by(
        User.created_at.desc()
    )

    users = db.scalars(
        statement
    ).unique().all()

    return [
        _user_to_list_item(user)
        for user in users
    ]


# ============================================================
# PENDING USERS
# ============================================================

def get_pending_users(
    db: Session,
):
    """
    Return all accounts waiting for activation.
    """

    return list_users(
        db=db,
        account_status="PENDING",
    )


# ============================================================
# GET USER DETAILS
# ============================================================

def get_user_details(
    db: Session,
    user_id: int,
):
    """
    Return complete details of one user.
    """

    user = db.scalar(
        select(User)
        .options(
            joinedload(User.role),
            joinedload(User.student_profile),
            joinedload(User.faculty_profile),
            joinedload(User.admin_profile),
        )
        .where(User.user_id == user_id)
    )

    if user is None:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User account not found.",
        )

    return _user_to_detail(user)


# ============================================================
# CHANGE ACCOUNT STATUS
# ============================================================

def change_account_status(
    db: Session,
    target_user_id: int,
    new_status: str,
    current_admin: User,
):
    """
    Change the account status of another user.

    The Admin cannot change their own account status.
    """

    normalized_status = new_status.strip().upper()

    allowed_statuses = {
        "PENDING",
        "ACTIVE",
        "SUSPENDED",
        "INACTIVE",
    }

    if normalized_status not in allowed_statuses:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid account status.",
        )

    # --------------------------------------------------------
    # PREVENT SELF STATUS CHANGE
    # --------------------------------------------------------

    if target_user_id == current_admin.user_id:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "An administrator cannot change "
                "their own account status."
            ),
        )

    # --------------------------------------------------------
    # FIND TARGET USER
    # --------------------------------------------------------

    user = db.scalar(
        select(User)
        .options(
            joinedload(User.role)
        )
        .where(User.user_id == target_user_id)
    )

    if user is None:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User account not found.",
        )

    previous_status = _normalize_status(
        user.account_status
    )

    # Faculty public registrations must be approved through the
    # dedicated Faculty-registration workflow so that requested
    # Semester + Section + Subject combinations are verified and
    # activated atomically with the account.
    target_role_name = ""

    if user.role is not None:
        target_role_name = str(user.role.role_name)

        if "." in target_role_name:
            target_role_name = target_role_name.split(".")[-1]

        target_role_name = target_role_name.upper()

    if (
        normalized_status == "ACTIVE"
        and previous_status == "PENDING"
        and target_role_name == "FACULTY"
    ):
        pending_faculty_request = db.execute(
            text(
                """
                SELECT far.request_id
                FROM faculty_assignment_requests far
                INNER JOIN faculty_profiles fp
                    ON fp.faculty_id = far.faculty_id
                WHERE fp.user_id = :user_id
                  AND far.request_status = 'PENDING'
                LIMIT 1
                """
            ),
            {"user_id": target_user_id},
        ).scalar()

        if pending_faculty_request is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "Pending Faculty registrations must be approved "
                    "through the Faculty registration review endpoint so "
                    "their requested academic assignments are verified."
                ),
            )

    if (
        normalized_status == "ACTIVE"
        and previous_status == "PENDING"
        and target_role_name == "STUDENT"
    ):
        pending_student_request = db.execute(
            text(
                """
                SELECT ssrr.request_id
                FROM student_section_registration_requests ssrr
                INNER JOIN student_profiles sp
                    ON sp.student_id = ssrr.student_id
                WHERE sp.user_id = :user_id
                  AND ssrr.request_status = 'PENDING'
                LIMIT 1
                """
            ),
            {"user_id": target_user_id},
        ).scalar()

        if pending_student_request is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "Pending Student registrations must be approved "
                    "through the Student registration review endpoint so "
                    "the requested Section is verified and assigned."
                ),
            )

    # --------------------------------------------------------
    # NO-OP CHECK
    # --------------------------------------------------------

    if previous_status == normalized_status:

        role_name = ""

        if user.role is not None:
            role_name = str(user.role.role_name)

            if "." in role_name:
                role_name = role_name.split(".")[-1]

            role_name = role_name.upper()

        return {
            "user_id": user.user_id,
            "email": user.email,
            "full_name": user.full_name,
            "role_name": role_name,
            "previous_status": previous_status,
            "account_status": normalized_status,
            "message": (
                f"Account is already {normalized_status}."
            ),
        }

    # --------------------------------------------------------
    # UPDATE STATUS
    # --------------------------------------------------------

    user.account_status = normalized_status

    try:

        db.commit()

        db.refresh(user)

    except Exception as exc:

        db.rollback()

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail={
                "status": "error",
                "message": "Unable to update account status.",
                "error": str(exc),
            },
        )

    # --------------------------------------------------------
    # ROLE
    # --------------------------------------------------------

    role_name = ""

    if user.role is not None:
        role_name = str(user.role.role_name)

        if "." in role_name:
            role_name = role_name.split(".")[-1]

        role_name = role_name.upper()

    # --------------------------------------------------------
    # SUCCESS MESSAGE
    # --------------------------------------------------------

    messages = {
        "ACTIVE": "Account activated successfully.",
        "SUSPENDED": "Account suspended successfully.",
        "INACTIVE": "Account deactivated successfully.",
        "PENDING": "Account moved back to pending status.",
    }

    return {
        "user_id": user.user_id,
        "email": user.email,
        "full_name": user.full_name,
        "role_name": role_name,
        "previous_status": previous_status,
        "account_status": normalized_status,
        "message": messages[normalized_status],
    }


# ============================================================
# ACADEMIC PROGRAM / COURSE-BRANCH MANAGEMENT
# ============================================================

def _program_row_to_dict(row, message: str | None = None) -> dict:
    result = {
        "program_id": int(row["program_id"]),
        "program_code": row["program_code"],
        "program_name": row["program_name"],
        "description": row["description"],
        "default_semesters": int(row["default_semesters"]),
        "is_active": bool(row["is_active"]),
    }
    if message is not None:
        result["message"] = message
    return result


def _get_program_row(db: Session, program_id: int):
    row = db.execute(
        text(
            """
            SELECT program_id, program_code, program_name, description, default_semesters, is_active
            FROM degree_programs
            WHERE program_id = :program_id
            LIMIT 1
            """
        ),
        {"program_id": program_id},
    ).mappings().first()
    if row is None:
        raise HTTPException(status_code=404, detail="Program not found.")
    return row


def list_programs(db: Session, active_only: bool = False) -> list[dict]:
    where_sql = " WHERE is_active = 1" if active_only else ""
    rows = db.execute(
        text(
            """
            SELECT program_id, program_code, program_name, description, default_semesters, is_active
            FROM degree_programs
            """ + where_sql + " ORDER BY program_name, program_id"
        )
    ).mappings().all()
    return [_program_row_to_dict(row) for row in rows]


def create_program(
    db: Session,
    program_code: str,
    program_name: str,
    description: str | None = None,
    default_semesters: int = 8,
) -> dict:
    code = program_code.strip().upper()
    name = " ".join(program_name.strip().split())
    desc = description.strip() if description and description.strip() else None

    if default_semesters < 1 or default_semesters > 8:
        raise HTTPException(
            status_code=400,
            detail="Default semesters must be between 1 and 8.",
        )

    duplicate = db.execute(
        text(
            """
            SELECT program_id
            FROM degree_programs
            WHERE UPPER(program_code) = :code
               OR UPPER(program_name) = UPPER(:name)
            LIMIT 1
            """
        ),
        {"code": code, "name": name},
    ).scalar()
    if duplicate is not None:
        raise HTTPException(status_code=409, detail="A program with this code or name already exists.")

    try:
        result = db.execute(
            text(
                """
                INSERT INTO degree_programs (
                    program_name,
                    program_code,
                    description,
                    default_semesters,
                    is_active
                ) VALUES (
                    :name,
                    :code,
                    :description,
                    :default_semesters,
                    1
                )
                """
            ),
            {
                "name": name,
                "code": code,
                "description": desc,
                "default_semesters": default_semesters,
            },
        )
        program_id = int(result.lastrowid)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Program could not be created because the code or name already exists.") from exc
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Unable to create program: {exc}") from exc

    return _program_row_to_dict(
        _get_program_row(db, program_id),
        message="Program added successfully.",
    )



def update_program(
    db: Session,
    program_id: int,
    program_code: str,
    program_name: str,
    description: str | None = None,
    default_semesters: int = 8,
) -> dict:
    """
    Edit an existing degree program without changing its primary key.

    Existing branches remain attached by program_id. Changing the Program's
    default semester count affects only the default used when creating/editing
    course/branch duration; it does not silently rewrite existing branches.
    """

    _get_program_row(db, program_id)

    code = program_code.strip().upper()
    name = " ".join(program_name.strip().split())
    desc = (
        description.strip()
        if description is not None and description.strip()
        else None
    )

    if not code:
        raise HTTPException(
            status_code=400,
            detail="Program code is required.",
        )

    if len(name) < 2:
        raise HTTPException(
            status_code=400,
            detail="Program name is required.",
        )

    if default_semesters < 1 or default_semesters > 8:
        raise HTTPException(
            status_code=400,
            detail="Default semesters must be between 1 and 8.",
        )

    duplicate = db.execute(
        text(
            """
            SELECT program_id
            FROM degree_programs
            WHERE program_id <> :program_id
              AND (
                    UPPER(program_code) = :code
                    OR UPPER(program_name) = UPPER(:name)
                  )
            LIMIT 1
            """
        ),
        {
            "program_id": program_id,
            "code": code,
            "name": name,
        },
    ).scalar()

    if duplicate is not None:
        raise HTTPException(
            status_code=409,
            detail=(
                "Another program with this code or name already exists."
            ),
        )

    try:
        db.execute(
            text(
                """
                UPDATE degree_programs
                SET
                    program_code = :code,
                    program_name = :name,
                    description = :description,
                    default_semesters = :default_semesters
                WHERE program_id = :program_id
                """
            ),
            {
                "program_id": program_id,
                "code": code,
                "name": name,
                "description": desc,
                "default_semesters": default_semesters,
            },
        )
        db.commit()

    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=(
                "Program could not be updated because the new code or "
                "name conflicts with existing data."
            ),
        ) from exc
    except Exception:
        db.rollback()
        raise

    return _program_row_to_dict(
        _get_program_row(db, program_id),
        message="Program updated successfully.",
    )


def set_program_active_status(db: Session, program_id: int, is_active: bool) -> dict:
    """
    Activate or deactivate a Program.

    Deactivation is cascade-style:
    Program -> all active Branches/Courses -> all active Subjects

    Historical rows are preserved; only is_active flags are changed.

    Safety:
    Deactivation is blocked while any ACTIVE attendance session is using
    a subject under this Program.
    """

    current = _get_program_row(db, program_id)
    desired = bool(is_active)

    if bool(current["is_active"]) == desired:
        return _program_row_to_dict(
            current,
            message=(
                "Program is already active."
                if desired
                else "Program is already inactive."
            ),
        )

    if desired:
        # Intentionally do not auto-reactivate child Branches/Subjects.
        # Some may have been manually deactivated before the Program itself
        # was disabled, so restoring all of them would be unsafe.
        try:
            db.execute(
                text(
                    """
                    UPDATE degree_programs
                    SET is_active = 1
                    WHERE program_id = :program_id
                    """
                ),
                {"program_id": program_id},
            )
            db.commit()
        except Exception:
            db.rollback()
            raise

        return _program_row_to_dict(
            _get_program_row(db, program_id),
            message=(
                "Program activated successfully. "
                "Courses/branches and subjects remain in their current "
                "status and can be activated separately."
            ),
        )

    active_session_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM attendance_sessions a
                INNER JOIN subjects s
                    ON s.subject_id = a.subject_id
                WHERE s.program_id = :program_id
                  AND a.session_status = 'ACTIVE'
                """
            ),
            {"program_id": program_id},
        ).scalar()
        or 0
    )

    if active_session_count:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Program cannot be deactivated while "
                f"{active_session_count} ACTIVE attendance session(s) "
                "are running under it. Complete or cancel those sessions "
                "first."
            ),
        )

    active_branch_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM branches
                WHERE program_id = :program_id
                  AND is_active = 1
                """
            ),
            {"program_id": program_id},
        ).scalar()
        or 0
    )

    active_subject_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM subjects
                WHERE program_id = :program_id
                  AND is_active = 1
                """
            ),
            {"program_id": program_id},
        ).scalar()
        or 0
    )

    try:
        db.execute(
            text(
                """
                UPDATE subjects
                SET is_active = 0
                WHERE program_id = :program_id
                  AND is_active = 1
                """
            ),
            {"program_id": program_id},
        )

        db.execute(
            text(
                """
                UPDATE branches
                SET is_active = 0
                WHERE program_id = :program_id
                  AND is_active = 1
                """
            ),
            {"program_id": program_id},
        )

        db.execute(
            text(
                """
                UPDATE degree_programs
                SET is_active = 0
                WHERE program_id = :program_id
                """
            ),
            {"program_id": program_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return _program_row_to_dict(
        _get_program_row(db, program_id),
        message=(
            "Program deactivated successfully. "
            f"{active_branch_count} active course/branch record(s) and "
            f"{active_subject_count} active subject record(s) were also "
            "deactivated. Historical data was preserved."
        ),
    )


def delete_program(db: Session, program_id: int) -> dict:
    current = _get_program_row(db, program_id)
    branch_count = int(db.execute(text("SELECT COUNT(*) FROM branches WHERE program_id = :program_id"), {"program_id": program_id}).scalar() or 0)
    subject_count = int(db.execute(text("SELECT COUNT(*) FROM subjects WHERE program_id = :program_id"), {"program_id": program_id}).scalar() or 0)
    if branch_count or subject_count:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Program cannot be permanently deleted because it is linked to "
                f"{branch_count} course/branch record(s) and {subject_count} subject record(s). "
                "Deactivate it instead."
            ),
        )
    try:
        db.execute(text("DELETE FROM degree_programs WHERE program_id = :program_id"), {"program_id": program_id})
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Program is referenced by existing academic records. Deactivate it instead.") from exc
    result = _program_row_to_dict(current)
    result["message"] = "Program deleted permanently."
    return result


def _branch_row_to_dict(row, message: str | None = None) -> dict:
    result = {
        "branch_id": int(row["branch_id"]),
        "program_id": int(row["program_id"]),
        "program_code": row["program_code"],
        "program_name": row["program_name"],
        "branch_code": row["branch_code"],
        "branch_name": row["branch_name"],
        "total_semesters": int(row["total_semesters"]),
        "is_active": bool(row["is_active"]),
    }
    if message is not None:
        result["message"] = message
    return result


def _get_branch_row(db: Session, branch_id: int):
    row = db.execute(
        text(
            """
            SELECT b.branch_id, b.program_id, b.branch_code, b.branch_name,
                   b.total_semesters, b.is_active,
                   p.program_code, p.program_name
            FROM branches b
            INNER JOIN degree_programs p ON p.program_id = b.program_id
            WHERE b.branch_id = :branch_id
            LIMIT 1
            """
        ),
        {"branch_id": branch_id},
    ).mappings().first()
    if row is None:
        raise HTTPException(status_code=404, detail="Course/branch not found.")
    return row


def list_branches(
    db: Session,
    program_id: int | None = None,
    active_only: bool = False,
) -> list[dict]:
    conditions = []
    params = {}
    if program_id is not None:
        conditions.append("b.program_id = :program_id")
        params["program_id"] = program_id
    if active_only:
        conditions.append("b.is_active = 1")
    where_sql = " WHERE " + " AND ".join(conditions) if conditions else ""
    rows = db.execute(
        text(
            """
            SELECT b.branch_id, b.program_id, b.branch_code, b.branch_name,
                   b.total_semesters, b.is_active,
                   p.program_code, p.program_name
            FROM branches b
            INNER JOIN degree_programs p ON p.program_id = b.program_id
            """ + where_sql + " ORDER BY p.program_name, b.branch_name, b.branch_id"
        ),
        params,
    ).mappings().all()
    return [_branch_row_to_dict(row) for row in rows]


def create_branch(
    db: Session,
    program_id: int,
    branch_code: str,
    branch_name: str,
    total_semesters: int,
) -> dict:
    program = _get_program_row(db, program_id)
    if not bool(program["is_active"]):
        raise HTTPException(status_code=409, detail="Selected program is inactive.")

    if total_semesters < 1 or total_semesters > 8:
        raise HTTPException(
            status_code=400,
            detail="Course/branch semesters must be between 1 and 8.",
        )

    code = branch_code.strip().upper()
    name = " ".join(branch_name.strip().split())
    duplicate = db.execute(
        text(
            """
            SELECT branch_id
            FROM branches
            WHERE program_id = :program_id
              AND (UPPER(branch_code) = :code OR UPPER(branch_name) = UPPER(:name))
            LIMIT 1
            """
        ),
        {"program_id": program_id, "code": code, "name": name},
    ).scalar()
    if duplicate is not None:
        raise HTTPException(status_code=409, detail="This program already has a course/branch with the same code or name.")

    try:
        result = db.execute(
            text(
                """
                INSERT INTO branches (
                    program_id,
                    branch_name,
                    branch_code,
                    total_semesters,
                    is_active
                )
                VALUES (
                    :program_id,
                    :name,
                    :code,
                    :total_semesters,
                    1
                )
                """
            ),
            {
                "program_id": program_id,
                "name": name,
                "code": code,
                "total_semesters": total_semesters,
            },
        )
        branch_id = int(result.lastrowid)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Course/branch could not be created because a duplicate value already exists.") from exc
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Unable to create course/branch: {exc}") from exc

    return _branch_row_to_dict(
        _get_branch_row(db, branch_id),
        message="Course/branch added successfully.",
    )



def update_branch(
    db: Session,
    branch_id: int,
    program_id: int,
    branch_code: str,
    branch_name: str,
    total_semesters: int,
) -> dict:
    """
    Edit an existing Course / Branch.

    Safe rules:
    - Code/name may be corrected while preserving branch_id.
    - Semester duration may increase freely.
    - Duration may not be reduced below semesters already used by Subjects
      or Sections in this Branch.
    - Moving to another Program is allowed only when the Branch has no
      linked academic/profile records.
    """

    current = _get_branch_row(db, branch_id)
    target_program = _get_program_row(db, program_id)

    if not bool(target_program["is_active"]):
        raise HTTPException(
            status_code=409,
            detail="Selected parent program is inactive.",
        )

    if total_semesters < 1 or total_semesters > 8:
        raise HTTPException(
            status_code=400,
            detail="Course/branch semesters must be between 1 and 8.",
        )

    code = branch_code.strip().upper()
    name = " ".join(branch_name.strip().split())

    if not code:
        raise HTTPException(
            status_code=400,
            detail="Course / branch code is required.",
        )

    if len(name) < 2:
        raise HTTPException(
            status_code=400,
            detail="Course / branch name is required.",
        )

    duplicate = db.execute(
        text(
            """
            SELECT branch_id
            FROM branches
            WHERE branch_id <> :branch_id
              AND program_id = :program_id
              AND (
                    UPPER(branch_code) = :code
                    OR UPPER(branch_name) = UPPER(:name)
                  )
            LIMIT 1
            """
        ),
        {
            "branch_id": branch_id,
            "program_id": program_id,
            "code": code,
            "name": name,
        },
    ).scalar()

    if duplicate is not None:
        raise HTTPException(
            status_code=409,
            detail=(
                "The selected Program already has another Course / Branch "
                "with this code or name."
            ),
        )

    current_program_id = int(current["program_id"])
    moving_program = current_program_id != program_id

    if moving_program:
        dependencies = {
            "subjects": _reference_count(
                db, "subjects", "branch_id", branch_id
            ),
            "students": _reference_count(
                db, "student_profiles", "branch_id", branch_id
            ),
            "faculty": _reference_count(
                db, "faculty_profiles", "branch_id", branch_id
            ),
            "sections": _reference_count(
                db, "sections", "branch_id", branch_id
            ),
            "faculty assignment requests": _reference_count(
                db, "faculty_assignment_requests", "branch_id", branch_id
            ),
        }

        used = [
            f"{count} {label}"
            for label, count in dependencies.items()
            if count
        ]

        if used:
            raise HTTPException(
                status_code=409,
                detail=(
                    "This Course / Branch cannot be moved to another Program "
                    "because it is already referenced by "
                    + ", ".join(used)
                    + ". You may still edit its code, name or semester count."
                ),
            )

    # Prevent a duration reduction that would hide semesters already used
    # by subject or section data in this Branch.
    highest_subject_semester = db.execute(
        text(
            """
            SELECT MAX(sem.semester_number)
            FROM subjects s
            INNER JOIN semesters sem
                ON sem.semester_id = s.semester_id
            WHERE s.branch_id = :branch_id
            """
        ),
        {"branch_id": branch_id},
    ).scalar()

    highest_section_semester = None
    if _column_exists(db, "sections", "branch_id"):
        highest_section_semester = db.execute(
            text(
                """
                SELECT MAX(sem.semester_number)
                FROM sections sec
                INNER JOIN semesters sem
                    ON sem.semester_id = sec.semester_id
                WHERE sec.branch_id = :branch_id
                """
            ),
            {"branch_id": branch_id},
        ).scalar()

    highest_used = max(
        int(highest_subject_semester or 0),
        int(highest_section_semester or 0),
    )

    if total_semesters < highest_used:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Semester count cannot be reduced to {total_semesters} "
                f"because this Course / Branch already contains academic "
                f"records in Semester {highest_used}."
            ),
        )

    try:
        db.execute(
            text(
                """
                UPDATE branches
                SET
                    program_id = :program_id,
                    branch_code = :code,
                    branch_name = :name,
                    total_semesters = :total_semesters
                WHERE branch_id = :branch_id
                """
            ),
            {
                "branch_id": branch_id,
                "program_id": program_id,
                "code": code,
                "name": name,
                "total_semesters": total_semesters,
            },
        )
        db.commit()

    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=(
                "Course / branch could not be updated because the new "
                "values conflict with existing academic data."
            ),
        ) from exc
    except Exception:
        db.rollback()
        raise

    return _branch_row_to_dict(
        _get_branch_row(db, branch_id),
        message="Course / branch updated successfully.",
    )


def set_branch_active_status(db: Session, branch_id: int, is_active: bool) -> dict:
    """
    Activate or deactivate one Course / Branch.

    Deactivation is cascade-style:
    Branch/Course -> all active Subjects

    Historical data remains untouched.

    Safety:
    Deactivation is blocked while an ACTIVE attendance session is using
    a subject under this Branch.
    """

    current = _get_branch_row(db, branch_id)
    desired = bool(is_active)

    if bool(current["is_active"]) == desired:
        return _branch_row_to_dict(
            current,
            message=(
                "Course/branch is already active."
                if desired
                else "Course/branch is already inactive."
            ),
        )

    if desired:
        program = _get_program_row(
            db,
            int(current["program_id"]),
        )

        if not bool(program["is_active"]):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Activate the parent program first.",
            )

        try:
            db.execute(
                text(
                    """
                    UPDATE branches
                    SET is_active = 1
                    WHERE branch_id = :branch_id
                    """
                ),
                {"branch_id": branch_id},
            )
            db.commit()
        except Exception:
            db.rollback()
            raise

        return _branch_row_to_dict(
            _get_branch_row(db, branch_id),
            message=(
                "Course/branch activated successfully. "
                "Subjects remain in their current status and can be "
                "activated separately."
            ),
        )

    active_session_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM attendance_sessions a
                INNER JOIN subjects s
                    ON s.subject_id = a.subject_id
                WHERE s.branch_id = :branch_id
                  AND a.session_status = 'ACTIVE'
                """
            ),
            {"branch_id": branch_id},
        ).scalar()
        or 0
    )

    if active_session_count:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Course/branch cannot be deactivated while "
                f"{active_session_count} ACTIVE attendance session(s) "
                "are running under it. Complete or cancel those sessions "
                "first."
            ),
        )

    active_subject_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM subjects
                WHERE branch_id = :branch_id
                  AND is_active = 1
                """
            ),
            {"branch_id": branch_id},
        ).scalar()
        or 0
    )

    try:
        db.execute(
            text(
                """
                UPDATE subjects
                SET is_active = 0
                WHERE branch_id = :branch_id
                  AND is_active = 1
                """
            ),
            {"branch_id": branch_id},
        )

        db.execute(
            text(
                """
                UPDATE branches
                SET is_active = 0
                WHERE branch_id = :branch_id
                """
            ),
            {"branch_id": branch_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return _branch_row_to_dict(
        _get_branch_row(db, branch_id),
        message=(
            "Course/branch deactivated successfully. "
            f"{active_subject_count} active subject record(s) were also "
            "deactivated. Historical data was preserved."
        ),
    )


def _column_exists(db: Session, table_name: str, column_name: str) -> bool:
    value = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_NAME = :table_name
              AND COLUMN_NAME = :column_name
            """
        ),
        {"table_name": table_name, "column_name": column_name},
    ).scalar()
    return bool(value)


def _reference_count(db: Session, table_name: str, column_name: str, value: int) -> int:
    if not _column_exists(db, table_name, column_name):
        return 0
    # table/column names below are internal constants, not user input.
    return int(db.execute(text(f"SELECT COUNT(*) FROM {table_name} WHERE {column_name} = :value"), {"value": value}).scalar() or 0)


def delete_branch(db: Session, branch_id: int) -> dict:
    current = _get_branch_row(db, branch_id)
    dependencies = {
        "subjects": _reference_count(db, "subjects", "branch_id", branch_id),
        "students": _reference_count(db, "student_profiles", "branch_id", branch_id),
        "faculty": _reference_count(db, "faculty_profiles", "branch_id", branch_id),
        "sections": _reference_count(db, "sections", "branch_id", branch_id),
    }
    total = sum(dependencies.values())
    if total:
        details = ", ".join(f"{count} {name}" for name, count in dependencies.items() if count)
        raise HTTPException(
            status_code=409,
            detail=f"Course/branch cannot be permanently deleted because it is referenced by {details}. Deactivate it instead.",
        )
    try:
        db.execute(text("DELETE FROM branches WHERE branch_id = :branch_id"), {"branch_id": branch_id})
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Course/branch is referenced by existing academic records. Deactivate it instead.") from exc
    result = _branch_row_to_dict(current)
    result["message"] = "Course/branch deleted permanently."
    return result


# ============================================================
# ACADEMIC SUBJECT MANAGEMENT
# ============================================================

def get_subject_management_options(db: Session) -> dict:
    """Return active Program -> Branch and Semester choices for Admin."""

    programs = db.execute(
        text(
            """
            SELECT
                program_id,
                program_code,
                program_name,
                default_semesters
            FROM degree_programs
            WHERE is_active = 1
            ORDER BY program_name, program_id
            """
        )
    ).mappings().all()

    branches = db.execute(
        text(
            """
            SELECT
                branch_id,
                program_id,
                branch_code,
                branch_name,
                total_semesters
            FROM branches
            WHERE is_active = 1
            ORDER BY program_id, branch_name, branch_id
            """
        )
    ).mappings().all()

    semesters = db.execute(
        text(
            """
            SELECT
                semester_id,
                semester_number,
                semester_name
            FROM semesters
            ORDER BY semester_number, semester_id
            """
        )
    ).mappings().all()

    return {
        "programs": [dict(row) for row in programs],
        "branches": [dict(row) for row in branches],
        "semesters": [dict(row) for row in semesters],
    }


def _subject_select_sql() -> str:
    return """
        SELECT
            s.subject_id,
            s.subject_code,
            s.subject_name,
            s.credits,
            s.semester_id,
            s.branch_id,
            s.is_active,

            sem.semester_number,
            sem.semester_name,

            b.branch_code,
            b.branch_name,
            s.program_id,

            p.program_code,
            p.program_name

        FROM subjects s
        INNER JOIN semesters sem
            ON sem.semester_id = s.semester_id
        INNER JOIN branches b
            ON b.branch_id = s.branch_id
        INNER JOIN degree_programs p
            ON p.program_id = s.program_id
    """


def _subject_row_to_dict(row, message: str | None = None) -> dict:
    result = {
        "subject_id": int(row["subject_id"]),
        "program_id": int(row["program_id"]),
        "program_code": row["program_code"],
        "program_name": row["program_name"],
        "branch_id": int(row["branch_id"]),
        "branch_code": row["branch_code"],
        "branch_name": row["branch_name"],
        "semester_id": int(row["semester_id"]),
        "semester_number": row["semester_number"],
        "semester_name": row["semester_name"],
        "subject_code": row["subject_code"],
        "subject_name": row["subject_name"],
        "credits": (
            float(row["credits"])
            if row["credits"] is not None
            else None
        ),
        "is_active": bool(row["is_active"]),
    }

    if message is not None:
        result["message"] = message

    return result


def _get_subject_row(
    db: Session,
    subject_id: int,
):
    row = db.execute(
        text(
            _subject_select_sql()
            + " WHERE s.subject_id = :subject_id LIMIT 1"
        ),
        {"subject_id": subject_id},
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Subject not found.",
        )

    return row


def list_subjects(
    db: Session,
    program_id: int | None = None,
    branch_id: int | None = None,
    semester_id: int | None = None,
    active_only: bool = False,
) -> list[dict]:
    conditions: list[str] = []
    params: dict = {}

    if program_id is not None:
        conditions.append("s.program_id = :program_id")
        params["program_id"] = program_id

    if branch_id is not None:
        conditions.append("s.branch_id = :branch_id")
        params["branch_id"] = branch_id

    if semester_id is not None:
        conditions.append("s.semester_id = :semester_id")
        params["semester_id"] = semester_id

    if active_only:
        conditions.append("s.is_active = 1")

    where_sql = ""
    if conditions:
        where_sql = " WHERE " + " AND ".join(conditions)

    rows = db.execute(
        text(
            _subject_select_sql()
            + where_sql
            + """
              ORDER BY
                p.program_name,
                b.branch_name,
                sem.semester_number,
                s.subject_code,
                s.subject_name
            """
        ),
        params,
    ).mappings().all()

    return [_subject_row_to_dict(row) for row in rows]


def create_subject(
    db: Session,
    program_id: int,
    branch_id: int,
    semester_id: int,
    subject_code: str,
    subject_name: str,
    credits: float | None = None,
) -> dict:
    """Create one active subject in Program -> Branch -> Semester."""

    academic_scope = db.execute(
        text(
            """
            SELECT
                b.branch_id,
                b.branch_code,
                b.branch_name,
                b.program_id,
                b.total_semesters,
                b.is_active AS branch_is_active,
                p.program_code,
                p.program_name,
                p.is_active AS program_is_active
            FROM branches b
            INNER JOIN degree_programs p
                ON p.program_id = b.program_id
            WHERE b.branch_id = :branch_id
            LIMIT 1
            """
        ),
        {"branch_id": branch_id},
    ).mappings().first()

    if academic_scope is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Selected branch/department does not exist.",
        )

    if int(academic_scope["program_id"]) != program_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Selected branch does not belong to the selected program.",
        )

    if not bool(academic_scope["program_is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected program is inactive.",
        )

    if not bool(academic_scope["branch_is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected branch/department is inactive.",
        )

    semester = db.execute(
        text(
            """
            SELECT
                semester_id,
                semester_number,
                semester_name
            FROM semesters
            WHERE semester_id = :semester_id
            LIMIT 1
            """
        ),
        {"semester_id": semester_id},
    ).mappings().first()

    if semester is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Selected semester does not exist.",
        )

    semester_number = int(semester["semester_number"])
    total_semesters = int(academic_scope["total_semesters"])

    if semester_number > total_semesters:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"Selected course/branch has only {total_semesters} semester(s). "
                f"Semester {semester_number} is outside its configured duration."
            ),
        )

    normalized_code = subject_code.strip().upper()
    normalized_name = " ".join(subject_name.strip().split())

    if not normalized_code:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Subject code is required.",
        )

    if not normalized_name:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Subject name is required.",
        )

    duplicate = db.execute(
        text(
            """
            SELECT subject_id
            FROM subjects
            WHERE UPPER(subject_code) = :subject_code
            LIMIT 1
            """
        ),
        {"subject_code": normalized_code},
    ).scalar()

    if duplicate is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A subject with this subject code already exists.",
        )

    try:
        result = db.execute(
            text(
                """
                INSERT INTO subjects (
                    program_id,
                    branch_id,
                    semester_id,
                    subject_code,
                    subject_name,
                    credits,
                    is_active
                )
                VALUES (
                    :program_id,
                    :branch_id,
                    :semester_id,
                    :subject_code,
                    :subject_name,
                    :credits,
                    1
                )
                """
            ),
            {
                "program_id": program_id,
                "branch_id": branch_id,
                "semester_id": semester_id,
                "subject_code": normalized_code,
                "subject_name": normalized_name,
                "credits": credits,
            },
        )

        subject_id = int(result.lastrowid)
        db.commit()

    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Subject could not be created because a unique academic value already exists.",
        ) from exc
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Unable to create subject: {exc}",
        ) from exc

    row = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    return _subject_row_to_dict(
        row,
        message="Subject added successfully.",
    )


# ============================================================
# SUBJECT CATALOG FILE IMPORT
# ============================================================

_SUBJECT_IMPORT_MAX_BYTES = 10 * 1024 * 1024
_SUBJECT_IMPORT_MAX_ROWS = 5000


def _normalize_subject_import_header(value) -> str:
    text_value = str(value or "").strip().lower()
    text_value = re.sub(r"[^a-z0-9]+", "_", text_value)
    return text_value.strip("_")


def _subject_import_pick(row: dict, *aliases: str):
    for alias in aliases:
        key = _normalize_subject_import_header(alias)
        if key in row:
            value = row[key]
            if value is not None and str(value).strip() != "":
                return value
    return None


def _subject_import_dict_from_headers(headers, values) -> dict:
    normalized_headers = [
        _normalize_subject_import_header(value)
        for value in headers
    ]

    result: dict = {}
    for index, key in enumerate(normalized_headers):
        if not key:
            continue
        result[key] = values[index] if index < len(values) else None
    return result


def _subject_import_header_is_usable(headers) -> bool:
    normalized = {
        _normalize_subject_import_header(value)
        for value in headers
        if value is not None
    }

    groups = (
        {"program", "program_code", "program_name"},
        {"branch", "branch_code", "branch_name", "course", "course_code", "course_name"},
        {"semester", "semester_number", "sem"},
        {"subject_code", "code"},
        {"subject_name", "subject"},
    )

    return all(bool(normalized.intersection(group)) for group in groups)


def _parse_subject_import_csv(file_bytes: bytes) -> list[dict]:
    try:
        decoded = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail="CSV file must use UTF-8 encoding.",
        ) from exc

    reader = csv.reader(io.StringIO(decoded))
    rows = list(reader)
    if not rows:
        return []

    headers = rows[0]
    if not _subject_import_header_is_usable(headers):
        raise HTTPException(
            status_code=400,
            detail=(
                "CSV header is invalid. Required columns are Program, Branch, "
                "Semester, Subject Code and Subject Name. Credits is optional."
            ),
        )

    parsed: list[dict] = []
    for excel_row_number, values in enumerate(rows[1:], start=2):
        if not any(str(value or "").strip() for value in values):
            continue
        parsed.append({
            "row_number": excel_row_number,
            "data": _subject_import_dict_from_headers(headers, values),
        })
    return parsed


def _parse_subject_import_xlsx(file_bytes: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "Excel import dependency is missing. Install it with: "
                "pip install openpyxl"
            ),
        ) from exc

    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read Excel file: {exc}",
        ) from exc

    if not rows:
        return []

    headers = list(rows[0])
    if not _subject_import_header_is_usable(headers):
        raise HTTPException(
            status_code=400,
            detail=(
                "Excel header is invalid. Required columns are Program, Branch, "
                "Semester, Subject Code and Subject Name. Credits is optional."
            ),
        )

    parsed: list[dict] = []
    for excel_row_number, values in enumerate(rows[1:], start=2):
        if not any(str(value or "").strip() for value in values):
            continue
        parsed.append({
            "row_number": excel_row_number,
            "data": _subject_import_dict_from_headers(headers, list(values)),
        })
    return parsed


def _parse_subject_import_delimited_text(text_value: str) -> list[dict]:
    lines = [line.strip() for line in text_value.splitlines() if line.strip()]
    if not lines:
        return []

    delimiters = ("|", "\t", ",", ";")
    for delimiter in delimiters:
        for header_index, line in enumerate(lines[:30]):
            headers = [part.strip() for part in line.split(delimiter)]
            if not _subject_import_header_is_usable(headers):
                continue

            parsed: list[dict] = []
            for index, data_line in enumerate(lines[header_index + 1 :], start=header_index + 2):
                values = [part.strip() for part in data_line.split(delimiter)]
                if len(values) < 5:
                    continue
                parsed.append({
                    "row_number": index,
                    "data": _subject_import_dict_from_headers(headers, values),
                })
            return parsed

    return []


def _parse_subject_import_pdf(file_bytes: bytes) -> list[dict]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "PDF import dependency is missing. Install it with: "
                "pip install pdfplumber"
            ),
        ) from exc

    parsed: list[dict] = []
    extracted_text: list[str] = []
    running_row_number = 2

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    extracted_text.append(page_text)

                for table in page.extract_tables() or []:
                    if not table:
                        continue

                    header_index = None
                    headers = None

                    for candidate_index, candidate in enumerate(table[:8]):
                        if candidate and _subject_import_header_is_usable(candidate):
                            header_index = candidate_index
                            headers = candidate
                            break

                    if headers is None or header_index is None:
                        continue

                    for values in table[header_index + 1 :]:
                        if not values or not any(str(value or "").strip() for value in values):
                            continue

                        # Repeated table headers on later pages are ignored.
                        if _subject_import_header_is_usable(values):
                            continue

                        parsed.append({
                            "row_number": running_row_number,
                            "data": _subject_import_dict_from_headers(headers, values),
                        })
                        running_row_number += 1
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read PDF file: {exc}",
        ) from exc

    if parsed:
        return parsed

    # Fallback for PDFs whose extracted text is already delimiter-separated.
    fallback = _parse_subject_import_delimited_text("\n".join(extracted_text))
    if fallback:
        return fallback

    raise HTTPException(
        status_code=400,
        detail=(
            "No usable subject table was found in the PDF. PDF import supports "
            "text-based tables, not scanned/image-only PDFs. Use CSV or XLSX for "
            "the most reliable import."
        ),
    )


def _parse_subject_import_file(file_name: str, file_bytes: bytes) -> list[dict]:
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if extension == "csv":
        rows = _parse_subject_import_csv(file_bytes)
    elif extension == "xlsx":
        rows = _parse_subject_import_xlsx(file_bytes)
    elif extension == "pdf":
        rows = _parse_subject_import_pdf(file_bytes)
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Upload CSV, XLSX or PDF.",
        )

    if len(rows) > _SUBJECT_IMPORT_MAX_ROWS:
        raise HTTPException(
            status_code=400,
            detail=f"Import is limited to {_SUBJECT_IMPORT_MAX_ROWS} subject rows per file.",
        )

    return rows


def _subject_import_semester_number(value) -> int | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        number = int(value)
        return number if float(value) == float(number) else None

    match = re.search(r"\d+", str(value))
    if not match:
        return None
    return int(match.group(0))


def _subject_import_credit(value) -> float | None:
    if value is None or str(value).strip() == "":
        return None

    try:
        number = float(str(value).strip())
    except ValueError:
        return None

    return number


def import_subject_catalog(
    db: Session,
    file_name: str,
    content_base64: str,
    dry_run: bool = True,
) -> dict:
    """
    Validate and optionally import an academic subject catalog.

    Program and Branch values in the file are identifiers used to resolve
    EXISTING academic structure. They may be either code or exact name.
    Missing/inactive Program or Branch values are reported as row errors;
    this endpoint never silently invents academic structure.
    """

    safe_name = file_name.strip()
    if not safe_name:
        raise HTTPException(status_code=400, detail="File name is required.")

    try:
        file_bytes = base64.b64decode(content_base64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Uploaded file payload is invalid.") from exc

    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    if len(file_bytes) > _SUBJECT_IMPORT_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Uploaded file must be 10 MB or smaller.")

    parsed_rows = _parse_subject_import_file(safe_name, file_bytes)
    if not parsed_rows:
        raise HTTPException(status_code=400, detail="The uploaded file contains no subject rows.")

    program_rows = db.execute(
        text(
            """
            SELECT program_id, program_code, program_name, is_active
            FROM degree_programs
            """
        )
    ).mappings().all()

    branch_rows = db.execute(
        text(
            """
            SELECT branch_id, program_id, branch_code, branch_name,
                   total_semesters, is_active
            FROM branches
            """
        )
    ).mappings().all()

    semester_rows = db.execute(
        text(
            """
            SELECT semester_id, semester_number, semester_name
            FROM semesters
            """
        )
    ).mappings().all()

    existing_subject_rows = db.execute(
        text(
            """
            SELECT subject_id, program_id, branch_id, semester_id,
                   subject_code, subject_name, credits
            FROM subjects
            """
        )
    ).mappings().all()

    def key(value) -> str:
        return " ".join(str(value or "").strip().split()).upper()

    programs_by_code = {key(row["program_code"]): row for row in program_rows}
    programs_by_name = {key(row["program_name"]): row for row in program_rows}

    branches_by_program: dict[int, dict[str, dict]] = {}
    for row in branch_rows:
        program_id = int(row["program_id"])
        bucket = branches_by_program.setdefault(program_id, {})
        bucket[f"CODE:{key(row['branch_code'])}"] = row
        bucket[f"NAME:{key(row['branch_name'])}"] = row

    semesters_by_number = {
        int(row["semester_number"]): row
        for row in semester_rows
        if row["semester_number"] is not None
    }

    existing_by_code = {
        key(row["subject_code"]): row
        for row in existing_subject_rows
    }

    errors: list[dict] = []
    preview_rows: list[dict] = []
    ready_rows: list[dict] = []
    skipped_existing = 0
    file_codes: dict[str, tuple] = {}

    for item in parsed_rows:
        row_number = int(item["row_number"])
        raw = {
            _normalize_subject_import_header(k): v
            for k, v in dict(item["data"]).items()
        }

        program_value = _subject_import_pick(
            raw,
            "program_code", "program", "program_name",
        )
        branch_value = _subject_import_pick(
            raw,
            "branch_code", "branch", "course_code", "course", "branch_name", "course_name",
        )
        semester_value = _subject_import_pick(
            raw,
            "semester_number", "semester", "sem",
        )
        subject_code_value = _subject_import_pick(
            raw,
            "subject_code", "code",
        )
        subject_name_value = _subject_import_pick(
            raw,
            "subject_name", "subject",
        )
        credits_value = _subject_import_pick(
            raw,
            "credits", "credit",
        )

        program_text = " ".join(str(program_value or "").strip().split())
        branch_text = " ".join(str(branch_value or "").strip().split())
        subject_code = str(subject_code_value or "").strip().upper()
        subject_name = " ".join(str(subject_name_value or "").strip().split())
        semester_number = _subject_import_semester_number(semester_value)
        credits = _subject_import_credit(credits_value)

        row_errors: list[str] = []

        if not program_text:
            row_errors.append("Program is required.")
        if not branch_text:
            row_errors.append("Branch/Course is required.")
        if semester_number is None:
            row_errors.append("Semester must contain a valid semester number.")
        if not subject_code:
            row_errors.append("Subject Code is required.")
        if not subject_name:
            row_errors.append("Subject Name is required.")
        if credits_value is not None and str(credits_value).strip() != "" and credits is None:
            row_errors.append("Credits must be numeric.")
        if credits is not None and not (0 <= credits <= 30):
            row_errors.append("Credits must be between 0 and 30.")

        program = None
        branch = None
        semester = None

        if program_text:
            program = programs_by_code.get(key(program_text)) or programs_by_name.get(key(program_text))
            if program is None:
                row_errors.append(f"Program '{program_text}' does not exist.")
            elif not bool(program["is_active"]):
                row_errors.append(f"Program '{program_text}' is inactive.")

        if branch_text and program is not None:
            bucket = branches_by_program.get(int(program["program_id"]), {})
            branch = bucket.get(f"CODE:{key(branch_text)}") or bucket.get(f"NAME:{key(branch_text)}")
            if branch is None:
                row_errors.append(
                    f"Branch/Course '{branch_text}' does not exist under Program '{program_text}'."
                )
            elif not bool(branch["is_active"]):
                row_errors.append(f"Branch/Course '{branch_text}' is inactive.")

        if semester_number is not None:
            semester = semesters_by_number.get(semester_number)
            if semester is None:
                row_errors.append(f"Semester {semester_number} does not exist in the semester master.")
            elif branch is not None and semester_number > int(branch["total_semesters"]):
                row_errors.append(
                    f"{branch['branch_code']} has only {int(branch['total_semesters'])} semester(s); "
                    f"Semester {semester_number} is outside its configured duration."
                )

        action = "ERROR"

        if not row_errors and program is not None and branch is not None and semester is not None:
            fingerprint = (
                int(program["program_id"]),
                int(branch["branch_id"]),
                int(semester["semester_id"]),
                subject_name.casefold(),
                credits,
            )

            duplicate_in_file = file_codes.get(key(subject_code))
            if duplicate_in_file is not None:
                if duplicate_in_file == fingerprint:
                    action = "SKIP_DUPLICATE_ROW"
                    skipped_existing += 1
                else:
                    row_errors.append(
                        f"Subject code '{subject_code}' appears more than once with conflicting data in the file."
                    )
            else:
                file_codes[key(subject_code)] = fingerprint

            if not row_errors and action != "SKIP_DUPLICATE_ROW":
                existing = existing_by_code.get(key(subject_code))
                if existing is not None:
                    existing_credits = float(existing["credits"]) if existing["credits"] is not None else None
                    same_existing = (
                        int(existing["program_id"]) == fingerprint[0]
                        and int(existing["branch_id"]) == fingerprint[1]
                        and int(existing["semester_id"]) == fingerprint[2]
                        and str(existing["subject_name"]).strip().casefold() == fingerprint[3]
                        and existing_credits == fingerprint[4]
                    )
                    if same_existing:
                        action = "SKIP_EXISTING"
                        skipped_existing += 1
                    else:
                        row_errors.append(
                            f"Subject code '{subject_code}' already exists with different academic details."
                        )
                else:
                    action = "CREATE"
                    ready_rows.append({
                        "row_number": row_number,
                        "program_id": int(program["program_id"]),
                        "branch_id": int(branch["branch_id"]),
                        "semester_id": int(semester["semester_id"]),
                        "program": str(program["program_code"]),
                        "branch": str(branch["branch_code"]),
                        "semester_number": int(semester_number),
                        "subject_code": subject_code,
                        "subject_name": subject_name,
                        "credits": credits,
                    })

        if row_errors:
            action = "ERROR"
            for message in row_errors:
                errors.append({"row_number": row_number, "message": message})

        if len(preview_rows) < 100:
            preview_rows.append({
                "row_number": row_number,
                "program": program_text,
                "branch": branch_text,
                "semester_number": semester_number,
                "subject_code": subject_code,
                "subject_name": subject_name,
                "credits": credits,
                "action": action,
            })

    total_credits_to_import = round(
        sum(float(row["credits"] or 0) for row in ready_rows),
        2,
    )

    if not dry_run and errors:
        first_errors = "; ".join(
            f"Row {item['row_number']}: {item['message']}"
            for item in errors[:5]
        )
        more = "" if len(errors) <= 5 else f" (+{len(errors) - 5} more errors)"
        raise HTTPException(
            status_code=400,
            detail=(
                "Import was cancelled because validation errors were found. "
                f"{first_errors}{more}"
            ),
        )

    imported = 0

    if not dry_run and ready_rows:
        try:
            for row in ready_rows:
                db.execute(
                    text(
                        """
                        INSERT INTO subjects (
                            program_id,
                            branch_id,
                            semester_id,
                            subject_code,
                            subject_name,
                            credits,
                            is_active
                        ) VALUES (
                            :program_id,
                            :branch_id,
                            :semester_id,
                            :subject_code,
                            :subject_name,
                            :credits,
                            1
                        )
                        """
                    ),
                    row,
                )
            db.commit()
            imported = len(ready_rows)
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(
                status_code=409,
                detail=(
                    "Import could not be completed because the database changed "
                    "or a duplicate subject code was created. Preview the file again."
                ),
            ) from exc
        except Exception as exc:
            db.rollback()
            raise HTTPException(status_code=500, detail=f"Unable to import subjects: {exc}") from exc

    if dry_run:
        if errors:
            message = f"Preview found {len(errors)} validation error(s). No data was changed."
        else:
            message = (
                f"Preview successful: {len(ready_rows)} new subject(s) ready to import; "
                f"{skipped_existing} duplicate/existing row(s) will be skipped."
            )
    else:
        message = (
            f"Imported {imported} subject(s) successfully. "
            f"{skipped_existing} duplicate/existing row(s) skipped."
        )

    return {
        "file_name": safe_name,
        "dry_run": bool(dry_run),
        "total_rows": len(parsed_rows),
        "ready_to_import": len(ready_rows),
        "skipped_existing": skipped_existing,
        "imported": imported,
        "error_count": len(errors),
        "total_credits_to_import": total_credits_to_import,
        "errors": errors[:100],
        "preview_rows": preview_rows,
        "message": message,
    }



def update_subject(
    db: Session,
    subject_id: int,
    program_id: int,
    branch_id: int,
    semester_id: int,
    subject_code: str,
    subject_name: str,
    credits: float | None = None,
) -> dict:
    """
    Edit an existing subject.

    Subject code, name and credits may be corrected while preserving the
    subject ID and historical references.

    Moving a subject to another Program / Branch / Semester is allowed only
    while the subject has never been referenced by Faculty assignments,
    Faculty registration/assignment requests, or attendance sessions.
    This prevents historical academic records from being silently changed.
    """

    current = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    academic_scope = db.execute(
        text(
            """
            SELECT
                b.branch_id,
                b.branch_code,
                b.branch_name,
                b.program_id,
                b.total_semesters,
                b.is_active AS branch_is_active,
                p.program_code,
                p.program_name,
                p.is_active AS program_is_active
            FROM branches b
            INNER JOIN degree_programs p
                ON p.program_id = b.program_id
            WHERE b.branch_id = :branch_id
            LIMIT 1
            """
        ),
        {"branch_id": branch_id},
    ).mappings().first()

    if academic_scope is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Selected branch/department does not exist.",
        )

    if int(academic_scope["program_id"]) != program_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Selected branch does not belong to the selected program.",
        )

    if not bool(academic_scope["program_is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected program is inactive.",
        )

    if not bool(academic_scope["branch_is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected branch/department is inactive.",
        )

    semester = db.execute(
        text(
            """
            SELECT
                semester_id,
                semester_number,
                semester_name
            FROM semesters
            WHERE semester_id = :semester_id
            LIMIT 1
            """
        ),
        {"semester_id": semester_id},
    ).mappings().first()

    if semester is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Selected semester does not exist.",
        )

    semester_number = int(semester["semester_number"])
    total_semesters = int(academic_scope["total_semesters"])

    if semester_number > total_semesters:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"Selected course/branch has only {total_semesters} semester(s). "
                f"Semester {semester_number} is outside its configured duration."
            ),
        )

    normalized_code = subject_code.strip().upper()
    normalized_name = " ".join(subject_name.strip().split())

    if not normalized_code:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Subject code is required.",
        )

    if not normalized_name:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Subject name is required.",
        )

    duplicate = db.execute(
        text(
            """
            SELECT subject_id
            FROM subjects
            WHERE UPPER(subject_code) = :subject_code
              AND subject_id <> :subject_id
            LIMIT 1
            """
        ),
        {
            "subject_code": normalized_code,
            "subject_id": subject_id,
        },
    ).scalar()

    if duplicate is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Another subject with this subject code already exists.",
        )

    scope_changed = (
        int(current["program_id"]) != program_id
        or int(current["branch_id"]) != branch_id
        or int(current["semester_id"]) != semester_id
    )

    if scope_changed:
        dependencies = {
            "faculty assignment": _reference_count(
                db,
                "faculty_subject_assignments",
                "subject_id",
                subject_id,
            ),
            "faculty registration/assignment request": _reference_count(
                db,
                "faculty_assignment_requests",
                "subject_id",
                subject_id,
            ),
            "attendance session/history": _reference_count(
                db,
                "attendance_sessions",
                "subject_id",
                subject_id,
            ),
        }

        used_by = [
            f"{count} {label}{'' if count == 1 else 's'}"
            for label, count in dependencies.items()
            if count
        ]

        if used_by:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "Program, branch or semester cannot be changed because "
                    "this subject is already referenced by "
                    + ", ".join(used_by)
                    + ". You may still edit its code, name or credits, or "
                    "deactivate it and create a new subject in the new scope."
                ),
            )

    try:
        db.execute(
            text(
                """
                UPDATE subjects
                SET
                    program_id = :program_id,
                    branch_id = :branch_id,
                    semester_id = :semester_id,
                    subject_code = :subject_code,
                    subject_name = :subject_name,
                    credits = :credits
                WHERE subject_id = :subject_id
                """
            ),
            {
                "subject_id": subject_id,
                "program_id": program_id,
                "branch_id": branch_id,
                "semester_id": semester_id,
                "subject_code": normalized_code,
                "subject_name": normalized_name,
                "credits": credits,
            },
        )
        db.commit()

    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Unable to update the subject because the new values conflict "
                "with existing academic data."
            ),
        ) from exc
    except Exception:
        db.rollback()
        raise

    updated = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    return _subject_row_to_dict(
        updated,
        message="Subject updated successfully.",
    )


def set_subject_active_status(
    db: Session,
    subject_id: int,
    is_active: bool,
) -> dict:
    """
    Activate or deactivate one subject without removing academic history.

    Deactivation is blocked while an ACTIVE attendance session is using
    the subject. Historical sessions and assignments remain intact.
    """

    current = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    desired = bool(is_active)

    if bool(current["is_active"]) == desired:
        return _subject_row_to_dict(
            current,
            message=(
                "Subject is already active."
                if desired
                else "Subject is already inactive."
            ),
        )

    if not desired:
        active_session = db.execute(
            text(
                """
                SELECT session_id
                FROM attendance_sessions
                WHERE subject_id = :subject_id
                  AND session_status = 'ACTIVE'
                LIMIT 1
                """
            ),
            {"subject_id": subject_id},
        ).scalar()

        if active_session is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "This subject is currently being used by an ACTIVE "
                    "attendance session. Complete or cancel that session "
                    "before deactivating the subject."
                ),
            )

    try:
        db.execute(
            text(
                """
                UPDATE subjects
                SET is_active = :is_active
                WHERE subject_id = :subject_id
                """
            ),
            {
                "subject_id": subject_id,
                "is_active": 1 if desired else 0,
            },
        )
        db.commit()

    except Exception:
        db.rollback()
        raise

    updated = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    return _subject_row_to_dict(
        updated,
        message=(
            "Subject activated successfully."
            if desired
            else "Subject deactivated successfully."
        ),
    )


def delete_subject(
    db: Session,
    subject_id: int,
) -> dict:
    """
    Permanently delete only a subject that has never been referenced.

    Any subject used by faculty assignments, faculty registration
    assignment requests, or attendance sessions is preserved so academic
    and attendance history remains valid. Such a subject should be
    deactivated instead.
    """

    current = _get_subject_row(
        db=db,
        subject_id=subject_id,
    )

    reference_checks = (
        (
            "faculty_subject_assignments",
            "assignment_id",
            "faculty assignment",
        ),
        (
            "faculty_assignment_requests",
            "request_id",
            "faculty registration/assignment request",
        ),
        (
            "attendance_sessions",
            "session_id",
            "attendance session/history",
        ),
    )

    for table_name, id_column, label in reference_checks:
        reference = db.execute(
            text(
                f"""
                SELECT {id_column}
                FROM {table_name}
                WHERE subject_id = :subject_id
                LIMIT 1
                """
            ),
            {"subject_id": subject_id},
        ).scalar()

        if reference is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    f"This subject is referenced by an existing {label} "
                    "and cannot be permanently deleted. "
                    "Deactivate the subject instead."
                ),
            )

    try:
        db.execute(
            text(
                """
                DELETE FROM subjects
                WHERE subject_id = :subject_id
                """
            ),
            {"subject_id": subject_id},
        )
        db.commit()

    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This subject is referenced by existing system records "
                "and cannot be permanently deleted. "
                "Deactivate the subject instead."
            ),
        )
    except Exception:
        db.rollback()
        raise

    return _subject_row_to_dict(
        current,
        message="Subject deleted successfully.",
    )


# ============================================================
# FACULTY ASSIGNMENT MANAGEMENT
# ============================================================

def get_faculty_assignment_options(
    db: Session,
    semester_id: int | None = None,
) -> dict:
    """
    Return values used by the Admin Faculty Assignment form.

    Sections and subjects are optionally limited to one semester so
    the Flutter form can use dependent Semester -> Section/Subject
    dropdowns without hard-coded academic data.
    """

    faculty_rows = db.execute(
        text(
            """
            SELECT
                fp.faculty_id,
                fp.employee_id,
                fp.department,
                fp.designation,
                u.full_name,
                u.email,
                u.account_status
            FROM faculty_profiles fp
            INNER JOIN users u
                ON u.user_id = fp.user_id
            INNER JOIN roles r
                ON r.role_id = u.role_id
            WHERE r.role_name = 'FACULTY'
              AND u.account_status = 'ACTIVE'
            ORDER BY u.full_name, fp.employee_id
            """
        )
    ).mappings().all()

    semester_rows = db.execute(
        text(
            """
            SELECT
                semester_id,
                semester_number,
                semester_name
            FROM semesters
            ORDER BY semester_number, semester_id
            """
        )
    ).mappings().all()

    section_params: dict = {}
    section_where = "sec.is_active = 1"

    subject_params: dict = {}
    subject_where = "s.is_active = 1"

    if semester_id is not None:
        section_where += " AND sec.semester_id = :semester_id"
        subject_where += " AND s.semester_id = :semester_id"

        section_params["semester_id"] = semester_id
        subject_params["semester_id"] = semester_id

    section_rows = db.execute(
        text(
            f"""
            SELECT
                sec.section_id,
                sec.section_name,
                sec.semester_id,
                sec.program_id,
                sec.branch_id,
                sec.academic_year_id
            FROM sections sec
            WHERE {section_where}
            ORDER BY sec.semester_id, sec.section_name
            """
        ),
        section_params,
    ).mappings().all()

    subject_rows = db.execute(
        text(
            f"""
            SELECT
                s.subject_id,
                s.subject_code,
                s.subject_name,
                s.credits,
                s.semester_id,
                s.program_id,
                s.branch_id
            FROM subjects s
            WHERE {subject_where}
            ORDER BY s.semester_id, s.subject_code, s.subject_name
            """
        ),
        subject_params,
    ).mappings().all()

    return {
        "faculties": [
            {
                "faculty_id": int(row["faculty_id"]),
                "employee_id": row["employee_id"],
                "full_name": row["full_name"],
                "email": row["email"],
                "department": row["department"],
                "designation": row["designation"],
                "account_status": _normalize_status(
                    row["account_status"]
                ),
            }
            for row in faculty_rows
        ],
        "semesters": [
            {
                "semester_id": int(row["semester_id"]),
                "semester_number": row["semester_number"],
                "semester_name": row["semester_name"],
            }
            for row in semester_rows
        ],
        "sections": [
            {
                "section_id": int(row["section_id"]),
                "section_name": row["section_name"],
                "semester_id": int(row["semester_id"]),
                "program_id": (
                    int(row["program_id"])
                    if row["program_id"] is not None
                    else None
                ),
                "branch_id": (
                    int(row["branch_id"])
                    if row["branch_id"] is not None
                    else None
                ),
                "academic_year_id": (
                    int(row["academic_year_id"])
                    if row["academic_year_id"] is not None
                    else None
                ),
            }
            for row in section_rows
        ],
        "subjects": [
            {
                "subject_id": int(row["subject_id"]),
                "subject_code": row["subject_code"],
                "subject_name": row["subject_name"],
                "credits": (
                    float(row["credits"])
                    if row["credits"] is not None
                    else None
                ),
                "semester_id": int(row["semester_id"]),
                "program_id": (
                    int(row["program_id"])
                    if row["program_id"] is not None
                    else None
                ),
                "branch_id": (
                    int(row["branch_id"])
                    if row["branch_id"] is not None
                    else None
                ),
            }
            for row in subject_rows
        ],
    }



def create_faculty_assignment_section(
    db: Session,
    source_section_id: int,
    section_name: str,
) -> dict:
    """
    Add another section using an existing section as the academic-scope
    template.

    The new section inherits Program, Branch, Academic Year and Semester
    from the selected existing section.
    """

    source = db.execute(
        text(
            """
            SELECT
                section_id,
                section_name,
                program_id,
                branch_id,
                academic_year_id,
                semester_id,
                is_active
            FROM sections
            WHERE section_id = :section_id
            LIMIT 1
            """
        ),
        {"section_id": source_section_id},
    ).mappings().first()

    if source is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Source section not found.",
        )

    if not bool(source["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Select an active section before adding another section.",
        )

    if source["semester_id"] is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="The selected section does not have a semester configured.",
        )

    if source["academic_year_id"] is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "The selected section does not have an academic year "
                "configured, so its scope cannot be copied safely."
            ),
        )

    normalized_name = " ".join(section_name.strip().split())

    if normalized_name.lower().startswith("section "):
        normalized_name = normalized_name[8:].strip()

    if not normalized_name:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Section name is required.",
        )

    normalized_name = normalized_name.upper()

    duplicate = db.execute(
        text(
            """
            SELECT
                section_id,
                is_active
            FROM sections
            WHERE UPPER(section_name) = :section_name
              AND program_id <=> :program_id
              AND branch_id <=> :branch_id
              AND academic_year_id <=> :academic_year_id
              AND semester_id = :semester_id
            ORDER BY section_id DESC
            LIMIT 1
            """
        ),
        {
            "section_name": normalized_name,
            "program_id": source["program_id"],
            "branch_id": source["branch_id"],
            "academic_year_id": source["academic_year_id"],
            "semester_id": source["semester_id"],
        },
    ).mappings().first()

    try:
        if duplicate is not None:
            section_id = int(duplicate["section_id"])

            if bool(duplicate["is_active"]):
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=(
                        f"Section {normalized_name} already exists in "
                        "this academic scope."
                    ),
                )

            db.execute(
                text(
                    """
                    UPDATE sections
                    SET is_active = 1
                    WHERE section_id = :section_id
                    """
                ),
                {"section_id": section_id},
            )

            message = (
                f"Section {normalized_name} reactivated successfully."
            )

        else:
            result = db.execute(
                text(
                    """
                    INSERT INTO sections (
                        section_name,
                        program_id,
                        branch_id,
                        academic_year_id,
                        semester_id,
                        is_active
                    )
                    VALUES (
                        :section_name,
                        :program_id,
                        :branch_id,
                        :academic_year_id,
                        :semester_id,
                        1
                    )
                    """
                ),
                {
                    "section_name": normalized_name,
                    "program_id": source["program_id"],
                    "branch_id": source["branch_id"],
                    "academic_year_id": source["academic_year_id"],
                    "semester_id": source["semester_id"],
                },
            )

            section_id = int(result.lastrowid)
            message = f"Section {normalized_name} added successfully."

        db.commit()

    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise

    created = db.execute(
        text(
            """
            SELECT
                section_id,
                section_name,
                program_id,
                branch_id,
                academic_year_id,
                semester_id,
                is_active
            FROM sections
            WHERE section_id = :section_id
            LIMIT 1
            """
        ),
        {"section_id": section_id},
    ).mappings().first()

    if created is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Section was saved but could not be reloaded.",
        )

    return {
        "section_id": int(created["section_id"]),
        "section_name": created["section_name"],
        "program_id": (
            int(created["program_id"])
            if created["program_id"] is not None
            else None
        ),
        "branch_id": (
            int(created["branch_id"])
            if created["branch_id"] is not None
            else None
        ),
        "academic_year_id": (
            int(created["academic_year_id"])
            if created["academic_year_id"] is not None
            else None
        ),
        "semester_id": int(created["semester_id"]),
        "is_active": bool(created["is_active"]),
        "message": message,
    }


def _faculty_assignment_select_sql() -> str:
    return """
        SELECT
            fsa.assignment_id,

            fp.faculty_id,
            fp.employee_id,
            fu.full_name AS faculty_name,

            sem.semester_id,
            sem.semester_number,
            sem.semester_name,

            sec.section_id,
            sec.section_name,

            s.subject_id,
            s.subject_code,
            s.subject_name,

            fsa.is_active

        FROM faculty_subject_assignments fsa

        INNER JOIN faculty_profiles fp
            ON fp.faculty_id = fsa.faculty_id

        INNER JOIN users fu
            ON fu.user_id = fp.user_id

        INNER JOIN sections sec
            ON sec.section_id = fsa.section_id

        INNER JOIN semesters sem
            ON sem.semester_id = sec.semester_id

        INNER JOIN subjects s
            ON s.subject_id = fsa.subject_id
    """


def _faculty_assignment_row_to_dict(
    row,
    message: str | None = None,
) -> dict:
    result = {
        "assignment_id": int(row["assignment_id"]),

        "faculty_id": int(row["faculty_id"]),
        "faculty_name": row["faculty_name"],
        "employee_id": row["employee_id"],

        "semester_id": int(row["semester_id"]),
        "semester_number": row["semester_number"],
        "semester_name": row["semester_name"],

        "section_id": int(row["section_id"]),
        "section_name": row["section_name"],

        "subject_id": int(row["subject_id"]),
        "subject_code": row["subject_code"],
        "subject_name": row["subject_name"],

        "is_active": bool(row["is_active"]),
    }

    if message is not None:
        result["message"] = message

    return result


def list_faculty_assignments(
    db: Session,
    faculty_id: int | None = None,
    semester_id: int | None = None,
    active_only: bool = False,
) -> list[dict]:
    conditions: list[str] = []
    params: dict = {}

    if faculty_id is not None:
        conditions.append(
            "fsa.faculty_id = :faculty_id"
        )
        params["faculty_id"] = faculty_id

    if semester_id is not None:
        conditions.append(
            "sec.semester_id = :semester_id"
        )
        params["semester_id"] = semester_id

    if active_only:
        conditions.append(
            "fsa.is_active = 1"
        )

    where_sql = ""

    if conditions:
        where_sql = (
            " WHERE "
            + " AND ".join(conditions)
        )

    rows = db.execute(
        text(
            _faculty_assignment_select_sql()
            + where_sql
            + """
              ORDER BY
                fu.full_name,
                sem.semester_number,
                sec.section_name,
                s.subject_code,
                fsa.assignment_id
            """
        ),
        params,
    ).mappings().all()

    return [
        _faculty_assignment_row_to_dict(row)
        for row in rows
    ]


def create_faculty_assignment(
    db: Session,
    faculty_id: int,
    semester_id: int,
    section_id: int,
    subject_id: int,
) -> dict:
    """
    Assign one active Semester + Section + Subject combination
    to one active Faculty member.

    Semester is validated through both the section and subject.
    The existing faculty_subject_assignments table remains the
    source of truth, so no attendance-history table is changed.
    """

    faculty = db.execute(
        text(
            """
            SELECT
                fp.faculty_id,
                fp.employee_id,
                u.full_name,
                u.account_status,
                r.role_name
            FROM faculty_profiles fp
            INNER JOIN users u
                ON u.user_id = fp.user_id
            INNER JOIN roles r
                ON r.role_id = u.role_id
            WHERE fp.faculty_id = :faculty_id
            LIMIT 1
            """
        ),
        {
            "faculty_id": faculty_id,
        },
    ).mappings().first()

    if faculty is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Faculty profile not found.",
        )

    if str(faculty["role_name"]).upper() != "FACULTY":
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Selected user is not a faculty member.",
        )

    if _normalize_status(
        faculty["account_status"]
    ) != "ACTIVE":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Faculty account must be ACTIVE before "
                "academic assignments can be created."
            ),
        )

    section = db.execute(
        text(
            """
            SELECT
                section_id,
                section_name,
                semester_id,
                program_id,
                branch_id,
                academic_year_id,
                is_active
            FROM sections
            WHERE section_id = :section_id
            LIMIT 1
            """
        ),
        {
            "section_id": section_id,
        },
    ).mappings().first()

    if section is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Section not found.",
        )

    if not bool(section["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected section is inactive.",
        )

    if int(section["semester_id"]) != semester_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Selected section does not belong to "
                "the selected semester."
            ),
        )

    if section["academic_year_id"] is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Selected section does not have an academic year "
                "configured."
            ),
        )

    academic_year_id = int(
        section["academic_year_id"]
    )

    subject = db.execute(
        text(
            """
            SELECT
                subject_id,
                subject_code,
                subject_name,
                semester_id,
                program_id,
                branch_id,
                is_active
            FROM subjects
            WHERE subject_id = :subject_id
            LIMIT 1
            """
        ),
        {
            "subject_id": subject_id,
        },
    ).mappings().first()

    if subject is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Subject not found.",
        )

    if not bool(subject["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Selected subject is inactive.",
        )

    if int(subject["semester_id"]) != semester_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Selected subject does not belong to "
                "the selected semester."
            ),
        )

    if (
        section["program_id"] is not None
        and subject["program_id"] is not None
        and int(section["program_id"]) != int(subject["program_id"])
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Selected subject belongs to a different Program "
                "than the selected Section."
            ),
        )

    if (
        section["branch_id"] is not None
        and subject["branch_id"] is not None
        and int(section["branch_id"]) != int(subject["branch_id"])
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Selected subject belongs to a different Branch/Course "
                "than the selected Section."
            ),
        )

    semester_exists = db.execute(
        text(
            """
            SELECT semester_id
            FROM semesters
            WHERE semester_id = :semester_id
            LIMIT 1
            """
        ),
        {
            "semester_id": semester_id,
        },
    ).scalar()

    if semester_exists is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Semester not found.",
        )

    existing = db.execute(
        text(
            """
            SELECT
                assignment_id,
                is_active
            FROM faculty_subject_assignments
            WHERE faculty_id = :faculty_id
              AND subject_id = :subject_id
              AND section_id = :section_id
              AND academic_year_id = :academic_year_id
            ORDER BY assignment_id DESC
            LIMIT 1
            """
        ),
        {
            "faculty_id": faculty_id,
            "subject_id": subject_id,
            "section_id": section_id,
            "academic_year_id": academic_year_id,
        },
    ).mappings().first()

    if (
        existing is not None
        and bool(existing["is_active"])
    ):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This faculty already has the selected "
                "semester, section and subject assignment."
            ),
        )

    try:
        if existing is not None:
            assignment_id = int(
                existing["assignment_id"]
            )

            db.execute(
                text(
                    """
                    UPDATE faculty_subject_assignments
                    SET is_active = 1
                    WHERE assignment_id = :assignment_id
                    """
                ),
                {
                    "assignment_id": assignment_id,
                },
            )

            action_message = (
                "Faculty assignment reactivated successfully."
            )
        else:
            result = db.execute(
                text(
                    """
                    INSERT INTO faculty_subject_assignments (
                        faculty_id,
                        subject_id,
                        section_id,
                        academic_year_id,
                        is_active
                    )
                    VALUES (
                        :faculty_id,
                        :subject_id,
                        :section_id,
                        :academic_year_id,
                        1
                    )
                    """
                ),
                {
                    "faculty_id": faculty_id,
                    "subject_id": subject_id,
                    "section_id": section_id,
                    "academic_year_id": academic_year_id,
                },
            )

            assignment_id = int(
                result.lastrowid
            )

            action_message = (
                "Faculty assignment created successfully."
            )

        db.commit()

    except Exception:
        db.rollback()
        raise

    row = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {
            "assignment_id": assignment_id,
        },
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "Faculty assignment was saved but "
                "could not be retrieved."
            ),
        )

    return _faculty_assignment_row_to_dict(
        row,
        message=action_message,
    )



def activate_faculty_assignment(
    db: Session,
    assignment_id: int,
) -> dict:
    """Reactivate one previously deactivated Faculty assignment safely."""

    row = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {"assignment_id": assignment_id},
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Faculty assignment not found.",
        )

    if bool(row["is_active"]):
        return _faculty_assignment_row_to_dict(
            row,
            message="Faculty assignment is already active.",
        )

    faculty_status = db.execute(
        text(
            """
            SELECT u.account_status
            FROM faculty_profiles fp
            INNER JOIN users u
                ON u.user_id = fp.user_id
            WHERE fp.faculty_id = :faculty_id
            LIMIT 1
            """
        ),
        {"faculty_id": int(row["faculty_id"])},
    ).scalar()

    if _normalize_status(faculty_status) != "ACTIVE":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Faculty account must be ACTIVE before this assignment can be activated.",
        )

    scope = db.execute(
        text(
            """
            SELECT
                sec.is_active AS section_active,
                s.is_active AS subject_active,
                p.is_active AS program_active,
                b.is_active AS branch_active
            FROM sections sec
            INNER JOIN subjects s
                ON s.subject_id = :subject_id
            INNER JOIN degree_programs p
                ON p.program_id = s.program_id
            INNER JOIN branches b
                ON b.branch_id = s.branch_id
               AND b.program_id = p.program_id
            WHERE sec.section_id = :section_id
            LIMIT 1
            """
        ),
        {
            "subject_id": int(row["subject_id"]),
            "section_id": int(row["section_id"]),
        },
    ).mappings().first()

    if scope is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="The assignment's academic scope is no longer available.",
        )

    if not bool(scope["program_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Activate the parent Program first.",
        )

    if not bool(scope["branch_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Activate the Course / Branch first.",
        )

    if not bool(scope["section_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Activate the Section first.",
        )

    if not bool(scope["subject_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Activate the Subject first.",
        )

    duplicate_active = db.execute(
        text(
            """
            SELECT other.assignment_id
            FROM faculty_subject_assignments target
            INNER JOIN faculty_subject_assignments other
                ON other.faculty_id = target.faculty_id
               AND other.subject_id = target.subject_id
               AND other.section_id = target.section_id
               AND other.academic_year_id <=> target.academic_year_id
            WHERE target.assignment_id = :assignment_id
              AND other.assignment_id <> target.assignment_id
              AND other.is_active = 1
            LIMIT 1
            """
        ),
        {"assignment_id": assignment_id},
    ).scalar()

    if duplicate_active is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="An equivalent Faculty assignment is already active.",
        )

    try:
        db.execute(
            text(
                """
                UPDATE faculty_subject_assignments
                SET is_active = 1
                WHERE assignment_id = :assignment_id
                """
            ),
            {"assignment_id": assignment_id},
        )
        db.commit()
    except Exception:
        db.rollback()
        raise

    updated_row = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {"assignment_id": assignment_id},
    ).mappings().first()

    if updated_row is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Faculty assignment was activated but could not be reloaded.",
        )

    return _faculty_assignment_row_to_dict(
        updated_row,
        message="Faculty assignment activated successfully.",
    )


def deactivate_faculty_assignment(
    db: Session,
    assignment_id: int,
) -> dict:
    row = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {
            "assignment_id": assignment_id,
        },
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Faculty assignment not found.",
        )

    if not bool(row["is_active"]):
        return _faculty_assignment_row_to_dict(
            row,
            message="Faculty assignment is already inactive.",
        )

    active_session = db.execute(
        text(
            """
            SELECT session_id
            FROM attendance_sessions
            WHERE faculty_id = :faculty_id
              AND subject_id = :subject_id
              AND section_id = :section_id
              AND session_status = 'ACTIVE'
            LIMIT 1
            """
        ),
        {
            "faculty_id": int(row["faculty_id"]),
            "subject_id": int(row["subject_id"]),
            "section_id": int(row["section_id"]),
        },
    ).scalar()

    if active_session is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This assignment has an ACTIVE attendance "
                "session. Complete or cancel that session "
                "before removing the assignment."
            ),
        )

    try:
        db.execute(
            text(
                """
                UPDATE faculty_subject_assignments
                SET is_active = 0
                WHERE assignment_id = :assignment_id
                """
            ),
            {
                "assignment_id": assignment_id,
            },
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    updated = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {
            "assignment_id": assignment_id,
        },
    ).mappings().first()

    if updated is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "Faculty assignment was updated but "
                "could not be retrieved."
            ),
        )

    return _faculty_assignment_row_to_dict(
        updated,
        message="Faculty assignment removed successfully.",
    )






def delete_faculty_assignment(
    db: Session,
    assignment_id: int,
) -> dict:
    """
    Permanently delete only an UNUSED and already INACTIVE assignment.

    If the Faculty/Subject/Section combination has ever been used by an
    attendance session, preserve the assignment history and keep it inactive.
    """

    row = db.execute(
        text(
            _faculty_assignment_select_sql()
            + """
              WHERE fsa.assignment_id = :assignment_id
              LIMIT 1
            """
        ),
        {"assignment_id": assignment_id},
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Faculty assignment not found.",
        )

    if bool(row["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Deactivate this Faculty assignment before deleting it permanently.",
        )

    session_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM attendance_sessions
                WHERE faculty_id = :faculty_id
                  AND subject_id = :subject_id
                  AND section_id = :section_id
                """
            ),
            {
                "faculty_id": int(row["faculty_id"]),
                "subject_id": int(row["subject_id"]),
                "section_id": int(row["section_id"]),
            },
        ).scalar()
        or 0
    )

    if session_count:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This Faculty assignment has attendance-session history and "
                "cannot be deleted permanently. Keep it deactivated instead."
            ),
        )

    try:
        db.execute(
            text(
                """
                DELETE FROM faculty_subject_assignments
                WHERE assignment_id = :assignment_id
                """
            ),
            {"assignment_id": assignment_id},
        )
        db.commit()
    except Exception:
        db.rollback()
        raise

    return _faculty_assignment_row_to_dict(
        row,
        message="Faculty assignment deleted permanently.",
    )


# ============================================================
# CLASSROOM MANAGEMENT
# ============================================================

def _classroom_row_to_dict(
    row,
    message: str | None = None,
) -> dict:
    result = {
        "classroom_id": int(row["classroom_id"]),
        "building_name": row["building_name"],
        "room_number": row["room_number"],
        "latitude": float(row["latitude"]),
        "longitude": float(row["longitude"]),
        "geofence_radius_meters": float(
            row["geofence_radius_meters"]
        ),
        "is_active": bool(row["is_active"]),
        "created_at": row["created_at"],
    }

    if message is not None:
        result["message"] = message

    return result


def _get_classroom_row(
    db: Session,
    classroom_id: int,
):
    row = db.execute(
        text(
            """
            SELECT
                classroom_id,
                building_name,
                room_number,
                latitude,
                longitude,
                geofence_radius_meters,
                is_active,
                created_at
            FROM classrooms
            WHERE classroom_id = :classroom_id
            LIMIT 1
            """
        ),
        {"classroom_id": classroom_id},
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Classroom not found.",
        )

    return row


def list_classrooms(
    db: Session,
    active_only: bool = False,
) -> list[dict]:
    where_sql = "WHERE is_active = 1" if active_only else ""

    rows = db.execute(
        text(
            f"""
            SELECT
                classroom_id,
                building_name,
                room_number,
                latitude,
                longitude,
                geofence_radius_meters,
                is_active,
                created_at
            FROM classrooms
            {where_sql}
            ORDER BY building_name, room_number, classroom_id
            """
        )
    ).mappings().all()

    return [
        _classroom_row_to_dict(row)
        for row in rows
    ]


def create_classroom(
    db: Session,
    data,
) -> dict:
    building_name = data.building_name.strip()
    room_number = data.room_number.strip()

    existing = db.execute(
        text(
            """
            SELECT classroom_id
            FROM classrooms
            WHERE building_name = :building_name
              AND room_number = :room_number
            LIMIT 1
            """
        ),
        {
            "building_name": building_name,
            "room_number": room_number,
        },
    ).scalar()

    if existing is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "A classroom with this building name and "
                "room number already exists."
            ),
        )

    try:
        result = db.execute(
            text(
                """
                INSERT INTO classrooms (
                    building_name,
                    room_number,
                    latitude,
                    longitude,
                    geofence_radius_meters,
                    is_active
                )
                VALUES (
                    :building_name,
                    :room_number,
                    :latitude,
                    :longitude,
                    :geofence_radius_meters,
                    :is_active
                )
                """
            ),
            {
                "building_name": building_name,
                "room_number": room_number,
                "latitude": data.latitude,
                "longitude": data.longitude,
                "geofence_radius_meters": data.geofence_radius_meters,
                "is_active": 1 if data.is_active else 0,
            },
        )

        classroom_id = int(result.lastrowid)
        db.commit()

    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "A classroom with this building name and "
                "room number already exists."
            ),
        )
    except Exception:
        db.rollback()
        raise

    row = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    return _classroom_row_to_dict(
        row,
        message="Classroom created successfully.",
    )


def update_classroom(
    db: Session,
    classroom_id: int,
    data,
) -> dict:
    current = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    values = data.model_dump(exclude_unset=True)

    if not values:
        return _classroom_row_to_dict(
            current,
            message="No classroom changes were provided.",
        )

    if "building_name" in values:
        values["building_name"] = values["building_name"].strip()

    if "room_number" in values:
        values["room_number"] = values["room_number"].strip()

    building_name = values.get(
        "building_name",
        current["building_name"],
    )
    room_number = values.get(
        "room_number",
        current["room_number"],
    )

    duplicate = db.execute(
        text(
            """
            SELECT classroom_id
            FROM classrooms
            WHERE building_name = :building_name
              AND room_number = :room_number
              AND classroom_id <> :classroom_id
            LIMIT 1
            """
        ),
        {
            "building_name": building_name,
            "room_number": room_number,
            "classroom_id": classroom_id,
        },
    ).scalar()

    if duplicate is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Another classroom already uses this building "
                "name and room number."
            ),
        )

    allowed_fields = {
        "building_name",
        "room_number",
        "latitude",
        "longitude",
        "geofence_radius_meters",
    }

    assignments = []
    params = {"classroom_id": classroom_id}

    for key, value in values.items():
        if key not in allowed_fields:
            continue
        assignments.append(f"{key} = :{key}")
        params[key] = value

    if not assignments:
        return _classroom_row_to_dict(
            current,
            message="No classroom changes were provided.",
        )

    try:
        db.execute(
            text(
                "UPDATE classrooms SET "
                + ", ".join(assignments)
                + " WHERE classroom_id = :classroom_id"
            ),
            params,
        )
        db.commit()

    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Another classroom already uses this building "
                "name and room number."
            ),
        )
    except Exception:
        db.rollback()
        raise

    updated = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    return _classroom_row_to_dict(
        updated,
        message="Classroom updated successfully.",
    )


def set_classroom_active_status(
    db: Session,
    classroom_id: int,
    is_active: bool,
) -> dict:
    current = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    desired = bool(is_active)

    if bool(current["is_active"]) == desired:
        return _classroom_row_to_dict(
            current,
            message=(
                "Classroom is already active."
                if desired
                else "Classroom is already inactive."
            ),
        )

    if not desired:
        active_session = db.execute(
            text(
                """
                SELECT session_id
                FROM attendance_sessions
                WHERE classroom_id = :classroom_id
                  AND session_status = 'ACTIVE'
                LIMIT 1
                """
            ),
            {"classroom_id": classroom_id},
        ).scalar()

        if active_session is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "This classroom has an ACTIVE attendance session. "
                    "Complete or cancel that session before "
                    "deactivating the classroom."
                ),
            )

    try:
        db.execute(
            text(
                """
                UPDATE classrooms
                SET is_active = :is_active
                WHERE classroom_id = :classroom_id
                """
            ),
            {
                "is_active": 1 if desired else 0,
                "classroom_id": classroom_id,
            },
        )
        db.commit()

    except Exception:
        db.rollback()
        raise

    updated = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    return _classroom_row_to_dict(
        updated,
        message=(
            "Classroom activated successfully."
            if desired
            else "Classroom deactivated successfully."
        ),
    )


def delete_classroom(
    db: Session,
    classroom_id: int,
) -> dict:
    """
    Permanently delete only a classroom that has never been used.

    Any classroom referenced by attendance_sessions is preserved so
    historical attendance and audit records remain intact. In that
    situation the Admin should deactivate the classroom instead.
    """
    current = _get_classroom_row(
        db=db,
        classroom_id=classroom_id,
    )

    attendance_reference = db.execute(
        text(
            """
            SELECT session_id
            FROM attendance_sessions
            WHERE classroom_id = :classroom_id
            LIMIT 1
            """
        ),
        {"classroom_id": classroom_id},
    ).scalar()

    if attendance_reference is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This classroom has attendance history or an attendance "
                "session and cannot be permanently deleted. "
                "Deactivate the classroom instead."
            ),
        )

    try:
        db.execute(
            text(
                """
                DELETE FROM classrooms
                WHERE classroom_id = :classroom_id
                """
            ),
            {"classroom_id": classroom_id},
        )
        db.commit()

    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This classroom is referenced by existing system records "
                "and cannot be permanently deleted. "
                "Deactivate the classroom instead."
            ),
        )
    except Exception:
        db.rollback()
        raise

    return _classroom_row_to_dict(
        current,
        message="Classroom deleted successfully.",
    )


# ============================================================
# CREATE ANOTHER ADMIN ACCOUNT
# ============================================================

def create_admin_account(
    db: Session,
    data,
) -> dict:
    """Create an ACTIVE administrator account from an Admin-only API."""

    email = str(data.email).strip().lower()
    employee_id = data.employee_id.strip()

    existing_email = db.scalar(
        select(User).where(User.email == email)
    )

    if existing_email is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="An account with this email already exists.",
        )

    existing_employee_id = db.execute(
        text(
            """
            SELECT admin_id
            FROM admin_profiles
            WHERE employee_id = :employee_id
            LIMIT 1
            """
        ),
        {"employee_id": employee_id},
    ).scalar()

    if existing_employee_id is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="An Admin with this employee_id already exists.",
        )

    admin_role = db.scalar(
        select(Role).where(Role.role_name == "ADMIN")
    )

    if admin_role is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="ADMIN role is not configured.",
        )

    try:
        password_hash = hash_password(data.password)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        )

    user = User(
        role_id=admin_role.role_id,
        email=email,
        password_hash=password_hash,
        full_name=data.full_name.strip(),
        phone=(data.phone.strip() if data.phone else None),
        account_status="ACTIVE",
    )

    try:
        db.add(user)
        db.flush()

        profile = AdminProfile(
            user_id=user.user_id,
            employee_id=employee_id,
            designation=(
                data.designation.strip()
                if data.designation
                else None
            ),
        )

        db.add(profile)
        db.commit()
        db.refresh(user)
        db.refresh(profile)

    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Unable to create Admin account: {exc}",
        )

    return {
        "user_id": int(user.user_id),
        "admin_id": int(profile.admin_id),
        "email": user.email,
        "full_name": user.full_name,
        "employee_id": profile.employee_id,
        "account_status": _normalize_status(user.account_status),
        "message": "Administrator account created successfully.",
    }


# ============================================================
# STUDENT REGISTRATION REVIEW
# ============================================================

def _validate_student_registration_section(
    db: Session,
    section_id: int,
) -> dict:
    section = db.execute(
        text(
            """
            SELECT
                sec.section_id,
                sec.section_name,
                sec.semester_id,
                sec.is_active,
                sem.semester_number,
                sem.semester_name
            FROM sections sec
            LEFT JOIN semesters sem
                ON sem.semester_id = sec.semester_id
            WHERE sec.section_id = :section_id
            LIMIT 1
            """
        ),
        {"section_id": section_id},
    ).mappings().first()

    if section is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Requested Student section no longer exists.",
        )

    if not bool(section["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Requested Student section is inactive.",
        )

    if section["semester_id"] is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Requested Student section has no semester configured.",
        )

    return dict(section)


def list_student_registration_requests(
    db: Session,
    request_status: str = "PENDING",
) -> list[dict]:
    normalized = request_status.strip().upper()

    if normalized not in {"PENDING", "APPROVED", "REJECTED", "ALL"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "request_status must be PENDING, APPROVED, REJECTED or ALL."
            ),
        )

    where_request = ""
    params: dict = {}

    if normalized != "ALL":
        where_request = " AND ssrr.request_status = :request_status"
        params["request_status"] = normalized

    rows = db.execute(
        text(
            f"""
            SELECT
                ssrr.request_id,
                u.user_id,
                sp.student_id,
                u.email,
                u.full_name,
                u.phone,
                sp.enrollment_number,
                sp.roll_number,
                sp.date_of_birth,
                u.account_status,
                u.created_at,

                ssrr.section_id,
                sec.section_name,
                sec.semester_id,
                sem.semester_number,
                sem.semester_name,

                ssrr.request_status,
                ssrr.requested_at,
                ssrr.reviewed_at,
                ssrr.reviewed_by_user_id,
                ssrr.review_note

            FROM student_section_registration_requests ssrr
            INNER JOIN student_profiles sp
                ON sp.student_id = ssrr.student_id
            INNER JOIN users u
                ON u.user_id = sp.user_id
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN sections sec
                ON sec.section_id = ssrr.section_id
            LEFT JOIN semesters sem
                ON sem.semester_id = sec.semester_id

            WHERE r.role_name = 'STUDENT'
              {where_request}

            ORDER BY u.created_at DESC, ssrr.request_id DESC
            """
        ),
        params,
    ).mappings().all()

    return [
        {
            "request_id": int(row["request_id"]),
            "user_id": int(row["user_id"]),
            "student_id": int(row["student_id"]),
            "email": row["email"],
            "full_name": row["full_name"],
            "phone": row["phone"],
            "enrollment_number": row["enrollment_number"],
            "roll_number": row["roll_number"],
            "date_of_birth": row["date_of_birth"],
            "account_status": _normalize_status(row["account_status"]),
            "created_at": row["created_at"],
            "section_id": int(row["section_id"]),
            "section_name": row["section_name"],
            "semester_id": int(row["semester_id"]),
            "semester_number": row["semester_number"],
            "semester_name": row["semester_name"],
            "request_status": str(row["request_status"]).upper(),
            "requested_at": row["requested_at"],
            "reviewed_at": row["reviewed_at"],
            "reviewed_by_user_id": row["reviewed_by_user_id"],
            "review_note": row["review_note"],
        }
        for row in rows
    ]


def approve_student_registration(
    db: Session,
    user_id: int,
    current_admin: User,
    note: str | None = None,
) -> dict:
    student = db.execute(
        text(
            """
            SELECT
                u.user_id,
                u.email,
                u.full_name,
                u.account_status,
                r.role_name,
                sp.student_id,
                ssrr.request_id,
                ssrr.section_id,
                ssrr.request_status
            FROM users u
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN student_profiles sp
                ON sp.user_id = u.user_id
            INNER JOIN student_section_registration_requests ssrr
                ON ssrr.student_id = sp.student_id
            WHERE u.user_id = :user_id
            LIMIT 1
            """
        ),
        {"user_id": user_id},
    ).mappings().first()

    if student is None or str(student["role_name"]).upper() != "STUDENT":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pending Student registration not found.",
        )

    if _normalize_status(student["account_status"]) != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Only a PENDING Student registration can be approved.",
        )

    if str(student["request_status"]).upper() != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Student Section request is not PENDING.",
        )

    student_id = int(student["student_id"])
    section_id = int(student["section_id"])
    section = _validate_student_registration_section(
        db=db,
        section_id=section_id,
    )

    try:
        # A student should have only one active current Section.
        db.execute(
            text(
                """
                UPDATE student_section_assignments
                SET is_active = 0
                WHERE student_id = :student_id
                  AND is_active = 1
                """
            ),
            {"student_id": student_id},
        )

        existing_assignment = db.execute(
            text(
                """
                SELECT assignment_id
                FROM student_section_assignments
                WHERE student_id = :student_id
                  AND section_id = :section_id
                LIMIT 1
                """
            ),
            {
                "student_id": student_id,
                "section_id": section_id,
            },
        ).scalar()

        if existing_assignment is None:
            db.execute(
                text(
                    """
                    INSERT INTO student_section_assignments (
                        student_id,
                        section_id,
                        is_active
                    )
                    VALUES (
                        :student_id,
                        :section_id,
                        1
                    )
                    """
                ),
                {
                    "student_id": student_id,
                    "section_id": section_id,
                },
            )
        else:
            db.execute(
                text(
                    """
                    UPDATE student_section_assignments
                    SET is_active = 1,
                        assigned_at = CURRENT_TIMESTAMP
                    WHERE assignment_id = :assignment_id
                    """
                ),
                {"assignment_id": int(existing_assignment)},
            )

        db.execute(
            text(
                """
                UPDATE student_section_registration_requests
                SET
                    request_status = 'APPROVED',
                    reviewed_at = CURRENT_TIMESTAMP,
                    reviewed_by_user_id = :reviewed_by_user_id,
                    review_note = :review_note
                WHERE request_id = :request_id
                  AND request_status = 'PENDING'
                """
            ),
            {
                "request_id": int(student["request_id"]),
                "reviewed_by_user_id": current_admin.user_id,
                "review_note": note.strip() if note else None,
            },
        )

        db.execute(
            text(
                """
                UPDATE users
                SET account_status = 'ACTIVE'
                WHERE user_id = :user_id
                """
            ),
            {"user_id": user_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "user_id": int(student["user_id"]),
        "student_id": student_id,
        "email": student["email"],
        "full_name": student["full_name"],
        "account_status": "ACTIVE",
        "section_id": section_id,
        "section_name": section["section_name"],
        "message": (
            "Student registration approved and Section assigned successfully."
        ),
    }


def reject_student_registration(
    db: Session,
    user_id: int,
    current_admin: User,
    note: str | None = None,
) -> dict:
    student = db.execute(
        text(
            """
            SELECT
                u.user_id,
                u.email,
                u.full_name,
                u.account_status,
                r.role_name,
                sp.student_id,
                ssrr.request_id,
                ssrr.section_id,
                ssrr.request_status,
                sec.section_name
            FROM users u
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN student_profiles sp
                ON sp.user_id = u.user_id
            INNER JOIN student_section_registration_requests ssrr
                ON ssrr.student_id = sp.student_id
            INNER JOIN sections sec
                ON sec.section_id = ssrr.section_id
            WHERE u.user_id = :user_id
            LIMIT 1
            """
        ),
        {"user_id": user_id},
    ).mappings().first()

    if student is None or str(student["role_name"]).upper() != "STUDENT":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pending Student registration not found.",
        )

    if _normalize_status(student["account_status"]) != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Only a PENDING Student registration can be rejected.",
        )

    if str(student["request_status"]).upper() != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Student Section request is not PENDING.",
        )

    try:
        db.execute(
            text(
                """
                UPDATE student_section_registration_requests
                SET
                    request_status = 'REJECTED',
                    reviewed_at = CURRENT_TIMESTAMP,
                    reviewed_by_user_id = :reviewed_by_user_id,
                    review_note = :review_note
                WHERE request_id = :request_id
                  AND request_status = 'PENDING'
                """
            ),
            {
                "request_id": int(student["request_id"]),
                "reviewed_by_user_id": current_admin.user_id,
                "review_note": note.strip() if note else None,
            },
        )

        db.execute(
            text(
                """
                UPDATE users
                SET account_status = 'INACTIVE'
                WHERE user_id = :user_id
                """
            ),
            {"user_id": user_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "user_id": int(student["user_id"]),
        "student_id": int(student["student_id"]),
        "email": student["email"],
        "full_name": student["full_name"],
        "account_status": "INACTIVE",
        "section_id": int(student["section_id"]),
        "section_name": student["section_name"],
        "message": "Student registration rejected.",
    }


# ============================================================
# FACULTY REGISTRATION REVIEW HELPERS
# ============================================================

def _validate_faculty_registration_request_scope(
    db: Session,
    semester_id: int,
    section_id: int,
    subject_id: int,
) -> int:
    """Validate one requested scope and return its academic_year_id."""

    section = db.execute(
        text(
            """
            SELECT
                section_id,
                semester_id,
                academic_year_id,
                is_active
            FROM sections
            WHERE section_id = :section_id
            LIMIT 1
            """
        ),
        {"section_id": section_id},
    ).mappings().first()

    if section is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A requested section no longer exists.",
        )

    if not bool(section["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A requested section is inactive.",
        )

    if int(section["semester_id"]) != semester_id:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "A requested section no longer belongs to the "
                "requested semester."
            ),
        )

    if section["academic_year_id"] is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A requested section has no academic year configured.",
        )

    subject = db.execute(
        text(
            """
            SELECT
                subject_id,
                semester_id,
                is_active
            FROM subjects
            WHERE subject_id = :subject_id
            LIMIT 1
            """
        ),
        {"subject_id": subject_id},
    ).mappings().first()

    if subject is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A requested subject no longer exists.",
        )

    if not bool(subject["is_active"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A requested subject is inactive.",
        )

    if int(subject["semester_id"]) != semester_id:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "A requested subject no longer belongs to the "
                "requested semester."
            ),
        )

    return int(section["academic_year_id"])


def list_faculty_registration_requests(
    db: Session,
    request_status: str = "PENDING",
) -> list[dict]:
    """Return Faculty registrations and their requested assignments."""

    normalized = request_status.strip().upper()

    if normalized not in {"PENDING", "APPROVED", "REJECTED", "ALL"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "request_status must be PENDING, APPROVED, REJECTED or ALL."
            ),
        )

    where_request = ""
    params: dict = {}

    if normalized != "ALL":
        where_request = " AND far.request_status = :request_status"
        params["request_status"] = normalized

    rows = db.execute(
        text(
            f"""
            SELECT
                u.user_id,
                u.email,
                u.full_name,
                u.phone,
                u.account_status,
                u.created_at,

                fp.faculty_id,
                fp.employee_id,
                fp.department,
                fp.designation,

                far.request_id,
                far.semester_id,
                sem.semester_number,
                sem.semester_name,
                far.section_id,
                sec.section_name,
                far.subject_id,
                s.subject_code,
                s.subject_name,
                far.request_status,
                far.requested_at,
                far.reviewed_at,
                far.reviewed_by_user_id,
                far.review_note

            FROM users u
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN faculty_profiles fp
                ON fp.user_id = u.user_id
            INNER JOIN faculty_assignment_requests far
                ON far.faculty_id = fp.faculty_id
            INNER JOIN semesters sem
                ON sem.semester_id = far.semester_id
            INNER JOIN sections sec
                ON sec.section_id = far.section_id
            INNER JOIN subjects s
                ON s.subject_id = far.subject_id

            WHERE r.role_name = 'FACULTY'
              {where_request}

            ORDER BY
                u.created_at DESC,
                u.user_id DESC,
                far.request_id ASC
            """
        ),
        params,
    ).mappings().all()

    grouped: dict[int, dict] = {}

    for row in rows:
        user_id = int(row["user_id"])

        item = grouped.get(user_id)
        if item is None:
            item = {
                "user_id": user_id,
                "faculty_id": int(row["faculty_id"]),
                "email": row["email"],
                "full_name": row["full_name"],
                "phone": row["phone"],
                "employee_id": row["employee_id"],
                "department": row["department"],
                "designation": row["designation"],
                "account_status": _normalize_status(row["account_status"]),
                "created_at": row["created_at"],
                "assignments": [],
            }
            grouped[user_id] = item

        item["assignments"].append(
            {
                "request_id": int(row["request_id"]),
                "semester_id": int(row["semester_id"]),
                "semester_number": row["semester_number"],
                "semester_name": row["semester_name"],
                "section_id": int(row["section_id"]),
                "section_name": row["section_name"],
                "subject_id": int(row["subject_id"]),
                "subject_code": row["subject_code"],
                "subject_name": row["subject_name"],
                "request_status": str(row["request_status"]),
                "requested_at": row["requested_at"],
                "reviewed_at": row["reviewed_at"],
                "reviewed_by_user_id": (
                    int(row["reviewed_by_user_id"])
                    if row["reviewed_by_user_id"] is not None
                    else None
                ),
                "review_note": row["review_note"],
            }
        )

    return list(grouped.values())


def approve_faculty_registration(
    db: Session,
    user_id: int,
    current_admin: User,
    note: str | None = None,
) -> dict:
    """
    Activate a pending Faculty account and atomically convert every
    PENDING registration assignment request into an active assignment.
    """

    faculty = db.execute(
        text(
            """
            SELECT
                u.user_id,
                u.email,
                u.full_name,
                u.account_status,
                r.role_name,
                fp.faculty_id
            FROM users u
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN faculty_profiles fp
                ON fp.user_id = u.user_id
            WHERE u.user_id = :user_id
            LIMIT 1
            """
        ),
        {"user_id": user_id},
    ).mappings().first()

    if faculty is None or str(faculty["role_name"]).upper() != "FACULTY":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pending Faculty registration not found.",
        )

    if _normalize_status(faculty["account_status"]) != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Only a PENDING Faculty registration can be approved.",
        )

    faculty_id = int(faculty["faculty_id"])

    requests = db.execute(
        text(
            """
            SELECT
                request_id,
                semester_id,
                section_id,
                subject_id
            FROM faculty_assignment_requests
            WHERE faculty_id = :faculty_id
              AND request_status = 'PENDING'
            ORDER BY request_id
            """
        ),
        {"faculty_id": faculty_id},
    ).mappings().all()

    if not requests:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This Faculty registration has no pending academic "
                "assignment requests to approve."
            ),
        )

    validated: list[tuple[dict, int]] = []

    # Validate every request before changing any database row.
    for request in requests:
        academic_year_id = _validate_faculty_registration_request_scope(
            db=db,
            semester_id=int(request["semester_id"]),
            section_id=int(request["section_id"]),
            subject_id=int(request["subject_id"]),
        )
        validated.append((request, academic_year_id))

    try:
        for request, academic_year_id in validated:
            existing = db.execute(
                text(
                    """
                    SELECT assignment_id
                    FROM faculty_subject_assignments
                    WHERE faculty_id = :faculty_id
                      AND subject_id = :subject_id
                      AND section_id = :section_id
                      AND academic_year_id = :academic_year_id
                    ORDER BY assignment_id DESC
                    LIMIT 1
                    """
                ),
                {
                    "faculty_id": faculty_id,
                    "subject_id": int(request["subject_id"]),
                    "section_id": int(request["section_id"]),
                    "academic_year_id": academic_year_id,
                },
            ).scalar()

            if existing is None:
                db.execute(
                    text(
                        """
                        INSERT INTO faculty_subject_assignments (
                            faculty_id,
                            subject_id,
                            section_id,
                            academic_year_id,
                            is_active
                        )
                        VALUES (
                            :faculty_id,
                            :subject_id,
                            :section_id,
                            :academic_year_id,
                            1
                        )
                        """
                    ),
                    {
                        "faculty_id": faculty_id,
                        "subject_id": int(request["subject_id"]),
                        "section_id": int(request["section_id"]),
                        "academic_year_id": academic_year_id,
                    },
                )
            else:
                db.execute(
                    text(
                        """
                        UPDATE faculty_subject_assignments
                        SET is_active = 1
                        WHERE assignment_id = :assignment_id
                        """
                    ),
                    {"assignment_id": int(existing)},
                )

        db.execute(
            text(
                """
                UPDATE faculty_assignment_requests
                SET
                    request_status = 'APPROVED',
                    reviewed_at = CURRENT_TIMESTAMP,
                    reviewed_by_user_id = :reviewed_by_user_id,
                    review_note = :review_note
                WHERE faculty_id = :faculty_id
                  AND request_status = 'PENDING'
                """
            ),
            {
                "faculty_id": faculty_id,
                "reviewed_by_user_id": current_admin.user_id,
                "review_note": note.strip() if note else None,
            },
        )

        db.execute(
            text(
                """
                UPDATE users
                SET account_status = 'ACTIVE'
                WHERE user_id = :user_id
                """
            ),
            {"user_id": user_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "user_id": int(faculty["user_id"]),
        "faculty_id": faculty_id,
        "email": faculty["email"],
        "full_name": faculty["full_name"],
        "account_status": "ACTIVE",
        "processed_assignments": len(validated),
        "message": (
            "Faculty registration approved and academic assignments "
            "activated successfully."
        ),
    }


def reject_faculty_registration(
    db: Session,
    user_id: int,
    current_admin: User,
    note: str | None = None,
) -> dict:
    """Reject a pending Faculty registration without deleting history."""

    faculty = db.execute(
        text(
            """
            SELECT
                u.user_id,
                u.email,
                u.full_name,
                u.account_status,
                r.role_name,
                fp.faculty_id
            FROM users u
            INNER JOIN roles r
                ON r.role_id = u.role_id
            INNER JOIN faculty_profiles fp
                ON fp.user_id = u.user_id
            WHERE u.user_id = :user_id
            LIMIT 1
            """
        ),
        {"user_id": user_id},
    ).mappings().first()

    if faculty is None or str(faculty["role_name"]).upper() != "FACULTY":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pending Faculty registration not found.",
        )

    if _normalize_status(faculty["account_status"]) != "PENDING":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Only a PENDING Faculty registration can be rejected.",
        )

    faculty_id = int(faculty["faculty_id"])

    try:
        result = db.execute(
            text(
                """
                UPDATE faculty_assignment_requests
                SET
                    request_status = 'REJECTED',
                    reviewed_at = CURRENT_TIMESTAMP,
                    reviewed_by_user_id = :reviewed_by_user_id,
                    review_note = :review_note
                WHERE faculty_id = :faculty_id
                  AND request_status = 'PENDING'
                """
            ),
            {
                "faculty_id": faculty_id,
                "reviewed_by_user_id": current_admin.user_id,
                "review_note": note.strip() if note else None,
            },
        )

        db.execute(
            text(
                """
                UPDATE users
                SET account_status = 'INACTIVE'
                WHERE user_id = :user_id
                """
            ),
            {"user_id": user_id},
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "user_id": int(faculty["user_id"]),
        "faculty_id": faculty_id,
        "email": faculty["email"],
        "full_name": faculty["full_name"],
        "account_status": "INACTIVE",
        "processed_assignments": int(result.rowcount or 0),
        "message": "Faculty registration rejected.",
    }


# ============================================================
# ADMIN STUDENT FACE MANAGEMENT
# ============================================================

def get_student_face_status_for_admin(
    db: Session,
    student_id: int,
) -> dict:
    """Return face metadata only. Never return the biometric embedding."""

    row = db.execute(
        text(
            """
            SELECT
                u.user_id,
                u.full_name,
                sp.student_id,
                sp.enrollment_number,

                sfp.face_profile_id,
                sfp.embedding_model,
                sfp.embedding_dimensions,
                sfp.is_verified,
                sfp.is_active,
                sfp.created_at AS face_created_at,
                sfp.updated_at AS face_updated_at

            FROM student_profiles sp
            INNER JOIN users u
                ON u.user_id = sp.user_id
            LEFT JOIN student_face_profiles sfp
                ON sfp.student_id = sp.student_id
            WHERE sp.student_id = :student_id
            LIMIT 1
            """
        ),
        {"student_id": student_id},
    ).mappings().first()

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student profile not found.",
        )

    registered = row["face_profile_id"] is not None

    return {
        "user_id": int(row["user_id"]),
        "student_id": int(row["student_id"]),
        "full_name": row["full_name"],
        "enrollment_number": row["enrollment_number"],
        "registered": registered,
        "face_profile_id": (
            int(row["face_profile_id"])
            if row["face_profile_id"] is not None
            else None
        ),
        "embedding_model": row["embedding_model"],
        "embedding_dimensions": (
            int(row["embedding_dimensions"])
            if row["embedding_dimensions"] is not None
            else None
        ),
        "is_verified": bool(row["is_verified"]) if registered else False,
        "is_active": bool(row["is_active"]) if registered else False,
        "created_at": row["face_created_at"],
        "updated_at": row["face_updated_at"],
    }


def delete_student_face_profile_for_admin(
    db: Session,
    student_id: int,
) -> dict:
    """
    Permanently remove only the student's biometric face profile.

    Student account, student profile and attendance history are untouched.
    """

    student_exists = db.execute(
        text(
            """
            SELECT student_id
            FROM student_profiles
            WHERE student_id = :student_id
            LIMIT 1
            """
        ),
        {"student_id": student_id},
    ).scalar()

    if student_exists is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student profile not found.",
        )

    face_profile_id = db.execute(
        text(
            """
            SELECT face_profile_id
            FROM student_face_profiles
            WHERE student_id = :student_id
            LIMIT 1
            """
        ),
        {"student_id": student_id},
    ).scalar()

    if face_profile_id is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student face profile is not registered.",
        )

    try:
        db.execute(
            text(
                """
                DELETE FROM student_face_profiles
                WHERE student_id = :student_id
                  AND face_profile_id = :face_profile_id
                """
            ),
            {
                "student_id": student_id,
                "face_profile_id": int(face_profile_id),
            },
        )
        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "student_id": student_id,
        "deleted_face_profile_id": int(face_profile_id),
        "registered": False,
        "message": (
            "Student face profile deleted successfully. "
            "The student may register a new face profile."
        ),
    }


# ============================================================
# ADMIN ATTENDANCE REPORT HELPERS
# ============================================================

def _attendance_percentage(
    present: int,
    total: int,
) -> float:
    if total <= 0:
        return 0.0

    return round(
        (present / total) * 100,
        2,
    )


def _normalize_report_status(
    value,
) -> str:
    if value is None:
        return "NOT_MARKED"

    normalized = str(value).upper()

    if "." in normalized:
        normalized = normalized.split(".")[-1]

    return normalized


def _build_admin_report_filters(
    semester_id: int | None,
    section_id: int | None,
    subject_id: int | None,
    faculty_id: int | None,
    date_from,
    date_to,
) -> tuple[str, dict]:
    conditions = [
        "ats.session_status <> 'CANCELLED'",
    ]

    params: dict = {}

    if semester_id is not None:
        conditions.append(
            "sec.semester_id = :semester_id"
        )
        params["semester_id"] = semester_id

    if section_id is not None:
        conditions.append(
            "ats.section_id = :section_id"
        )
        params["section_id"] = section_id

    if subject_id is not None:
        conditions.append(
            "ats.subject_id = :subject_id"
        )
        params["subject_id"] = subject_id

    if faculty_id is not None:
        conditions.append(
            "ats.faculty_id = :faculty_id"
        )
        params["faculty_id"] = faculty_id

    if date_from is not None:
        conditions.append(
            "DATE(ats.started_at) >= :date_from"
        )
        params["date_from"] = date_from

    if date_to is not None:
        conditions.append(
            "DATE(ats.started_at) <= :date_to"
        )
        params["date_to"] = date_to

    return (
        "\n              AND ".join(
            conditions
        ),
        params,
    )


def _get_admin_report_rows(
    db: Session,
    semester_id: int | None = None,
    section_id: int | None = None,
    subject_id: int | None = None,
    faculty_id: int | None = None,
    date_from=None,
    date_to=None,
) -> list[dict]:
    where_sql, params = (
        _build_admin_report_filters(
            semester_id=semester_id,
            section_id=section_id,
            subject_id=subject_id,
            faculty_id=faculty_id,
            date_from=date_from,
            date_to=date_to,
        )
    )

    rows = db.execute(
        text(
            f"""
            SELECT
                ats.session_id,

                ats.faculty_id,
                fp.employee_id AS faculty_employee_id,
                fu.full_name AS faculty_name,
                fu.email AS faculty_email,
                fp.department AS faculty_department,
                fp.designation AS faculty_designation,

                ats.subject_id,
                s.subject_code,
                s.subject_name,

                ats.section_id,
                sec.section_name,
                sec.semester_id,

                ats.classroom_id,
                c.building_name,
                c.room_number,

                ats.started_at,
                ats.ended_at,
                ats.session_status,

                sp.student_id,
                sp.enrollment_number,
                sp.roll_number,
                su.full_name AS student_name,
                su.email AS student_email,

                ar.attendance_id,
                ar.marked_at,
                ar.distance_from_classroom_meters,
                ar.qr_verified,
                ar.geofence_verified,
                ar.face_verified,
                ar.liveness_verified,
                ar.attendance_status,
                ar.rejection_reason

            FROM attendance_sessions ats

            INNER JOIN faculty_profiles fp
                ON fp.faculty_id = ats.faculty_id

            INNER JOIN users fu
                ON fu.user_id = fp.user_id

            INNER JOIN subjects s
                ON s.subject_id = ats.subject_id

            INNER JOIN sections sec
                ON sec.section_id = ats.section_id

            LEFT JOIN classrooms c
                ON c.classroom_id = ats.classroom_id

            INNER JOIN student_section_assignments ssa
                ON ssa.section_id = ats.section_id
               AND ssa.is_active = 1

            INNER JOIN student_profiles sp
                ON sp.student_id = ssa.student_id

            INNER JOIN users su
                ON su.user_id = sp.user_id

            LEFT JOIN attendance_records ar
                ON ar.session_id = ats.session_id
               AND ar.student_id = sp.student_id

            WHERE {where_sql}

            ORDER BY
                ats.started_at DESC,
                ats.session_id DESC,
                sp.roll_number,
                su.full_name
            """
        ),
        params,
    ).mappings().all()

    result: list[dict] = []

    for row in rows:
        attendance_id = (
            int(row["attendance_id"])
            if row["attendance_id"] is not None
            else None
        )

        started_at = row["started_at"]
        ended_at = row["ended_at"]
        marked_at = row["marked_at"]

        attendance_status = (
            _normalize_report_status(
                row["attendance_status"]
            )
        )

        session_status = (
            _normalize_report_status(
                row["session_status"]
            )
        )

        result.append(
            {
                "session_id":
                    int(row["session_id"]),

                "faculty_id":
                    int(row["faculty_id"]),

                "faculty_employee_id":
                    row["faculty_employee_id"],

                "faculty_name":
                    row["faculty_name"],

                "faculty_email":
                    row["faculty_email"],

                "faculty_department":
                    row["faculty_department"],

                "faculty_designation":
                    row["faculty_designation"],

                "subject_id":
                    int(row["subject_id"]),

                "subject_code":
                    row["subject_code"],

                "subject_name":
                    row["subject_name"],

                "section_id":
                    int(row["section_id"]),

                "section_name":
                    row["section_name"],

                "semester_id":
                    int(row["semester_id"]),

                "classroom_id": (
                    int(row["classroom_id"])
                    if row["classroom_id"] is not None
                    else None
                ),

                "building_name":
                    row["building_name"],

                "room_number":
                    row["room_number"],

                "date": (
                    started_at.date().isoformat()
                    if started_at is not None
                    else None
                ),

                "started_at": (
                    started_at.isoformat()
                    if started_at is not None
                    else None
                ),

                "ended_at": (
                    ended_at.isoformat()
                    if ended_at is not None
                    else None
                ),

                "session_status":
                    session_status,

                "student_id":
                    int(row["student_id"]),

                "enrollment_number":
                    row["enrollment_number"],

                "roll_number":
                    row["roll_number"],

                "student_name":
                    row["student_name"],

                "student_email":
                    row["student_email"],

                "attendance_id":
                    attendance_id,

                "marked_at": (
                    marked_at.isoformat()
                    if marked_at is not None
                    else None
                ),

                "attendance_status":
                    attendance_status,

                "qr_verified": (
                    bool(row["qr_verified"])
                    if attendance_id is not None
                    else False
                ),

                "geofence_verified": (
                    bool(row["geofence_verified"])
                    if attendance_id is not None
                    else False
                ),

                "face_verified": (
                    bool(row["face_verified"])
                    if attendance_id is not None
                    else False
                ),

                "liveness_verified": (
                    bool(row["liveness_verified"])
                    if attendance_id is not None
                    else False
                ),

                "distance_from_classroom_meters": (
                    float(
                        row[
                            "distance_from_classroom_meters"
                        ]
                    )
                    if row[
                        "distance_from_classroom_meters"
                    ] is not None
                    else None
                ),

                "rejection_reason":
                    row["rejection_reason"],
            }
        )

    return result


def _admin_report_summary(
    rows: list[dict],
) -> dict:
    session_ids = {
        row["session_id"]
        for row in rows
    }

    student_ids = {
        row["student_id"]
        for row in rows
    }

    faculty_ids = {
        row["faculty_id"]
        for row in rows
    }

    total_opportunities = len(rows)

    attendance_records = sum(
        1
        for row in rows
        if row["attendance_id"] is not None
    )

    present = sum(
        1
        for row in rows
        if row["attendance_status"] == "PRESENT"
    )

    pending = sum(
        1
        for row in rows
        if row["attendance_status"] == "PENDING"
    )

    rejected = sum(
        1
        for row in rows
        if row["attendance_status"] == "REJECTED"
    )

    not_marked = sum(
        1
        for row in rows
        if row["attendance_status"] == "NOT_MARKED"
    )

    return {
        "total_sessions": len(session_ids),
        "total_students": len(student_ids),
        "total_faculty": len(faculty_ids),
        "attendance_opportunities":
            total_opportunities,
        "attendance_records":
            attendance_records,
        "present": present,
        "pending": pending,
        "rejected": rejected,
        "not_marked": not_marked,
        "attendance_percentage":
            _attendance_percentage(
                present=present,
                total=total_opportunities,
            ),
    }


def _group_admin_report_rows(
    rows: list[dict],
    key_fields: tuple[str, ...],
) -> list[dict]:
    groups: dict[tuple, dict] = {}

    for row in rows:
        key = tuple(
            row[field]
            for field in key_fields
        )

        group = groups.get(key)

        if group is None:
            group = {
                field: row[field]
                for field in key_fields
            }

            group.update(
                {
                    "session_ids": set(),
                    "student_ids": set(),
                    "attendance_records": 0,
                    "present": 0,
                    "pending": 0,
                    "rejected": 0,
                    "not_marked": 0,
                }
            )

            groups[key] = group

        group["session_ids"].add(
            row["session_id"]
        )

        group["student_ids"].add(
            row["student_id"]
        )

        if row["attendance_id"] is not None:
            group["attendance_records"] += 1

        status_name = row["attendance_status"]

        if status_name == "PRESENT":
            group["present"] += 1
        elif status_name == "PENDING":
            group["pending"] += 1
        elif status_name == "REJECTED":
            group["rejected"] += 1
        else:
            group["not_marked"] += 1

    result: list[dict] = []

    for group in groups.values():
        total = (
            group["present"]
            + group["pending"]
            + group["rejected"]
            + group["not_marked"]
        )

        # Count unique sessions and unique students for this
        # exact report group. Keep this explicit because these values
        # are displayed directly in the Admin grouped-report tables.
        session_ids = group.pop("session_ids")
        student_ids = group.pop("student_ids")

        group["session_count"] = len(
            {
                session_id
                for session_id in session_ids
                if session_id is not None
            }
        )

        group["student_count"] = len(
            {
                student_id
                for student_id in student_ids
                if student_id is not None
            }
        )

        group["attendance_opportunities"] = (
            total
        )

        group["attendance_percentage"] = (
            _attendance_percentage(
                present=group["present"],
                total=total,
            )
        )

        result.append(group)

    return result


def _admin_faculty_wise_report(
    rows: list[dict],
) -> list[dict]:
    result = _group_admin_report_rows(
        rows=rows,
        key_fields=(
            "faculty_id",
            "faculty_employee_id",
            "faculty_name",
            "faculty_department",
            "faculty_designation",
        ),
    )

    result.sort(
        key=lambda item: (
            str(item["faculty_name"] or ""),
            str(
                item["faculty_employee_id"]
                or ""
            ),
        )
    )

    return result


def _admin_subject_wise_report(
    rows: list[dict],
) -> list[dict]:
    result = _group_admin_report_rows(
        rows=rows,
        key_fields=(
            "subject_id",
            "subject_code",
            "subject_name",
        ),
    )

    result.sort(
        key=lambda item: (
            str(item["subject_code"] or ""),
            str(item["subject_name"] or ""),
        )
    )

    return result


def _admin_section_wise_report(
    rows: list[dict],
) -> list[dict]:
    result = _group_admin_report_rows(
        rows=rows,
        key_fields=(
            "semester_id",
            "section_id",
            "section_name",
        ),
    )

    result.sort(
        key=lambda item: (
            int(item["semester_id"] or 0),
            str(item["section_name"] or ""),
        )
    )

    return result


def _admin_day_wise_report(
    rows: list[dict],
) -> list[dict]:
    result = _group_admin_report_rows(
        rows=rows,
        key_fields=(
            "date",
        ),
    )

    result.sort(
        key=lambda item: (
            item["date"] or ""
        ),
        reverse=True,
    )

    return result


def _admin_session_wise_report(
    rows: list[dict],
) -> list[dict]:
    groups: dict[int, dict] = {}

    for row in rows:
        session_id = row["session_id"]

        group = groups.get(session_id)

        if group is None:
            group = {
                "session_id":
                    session_id,

                "date":
                    row["date"],

                "started_at":
                    row["started_at"],

                "ended_at":
                    row["ended_at"],

                "session_status":
                    row["session_status"],

                "faculty_id":
                    row["faculty_id"],

                "faculty_employee_id":
                    row["faculty_employee_id"],

                "faculty_name":
                    row["faculty_name"],

                "subject_id":
                    row["subject_id"],

                "subject_code":
                    row["subject_code"],

                "subject_name":
                    row["subject_name"],

                "section_id":
                    row["section_id"],

                "section_name":
                    row["section_name"],

                "semester_id":
                    row["semester_id"],

                "classroom_id":
                    row["classroom_id"],

                "building_name":
                    row["building_name"],

                "room_number":
                    row["room_number"],

                "expected_students": 0,
                "attendance_records": 0,
                "present": 0,
                "pending": 0,
                "rejected": 0,
                "not_marked": 0,
            }

            groups[session_id] = group

        group["expected_students"] += 1

        if row["attendance_id"] is not None:
            group["attendance_records"] += 1

        status_name = row["attendance_status"]

        if status_name == "PRESENT":
            group["present"] += 1
        elif status_name == "PENDING":
            group["pending"] += 1
        elif status_name == "REJECTED":
            group["rejected"] += 1
        else:
            group["not_marked"] += 1

    result = list(groups.values())

    for group in result:
        group["attendance_percentage"] = (
            _attendance_percentage(
                present=group["present"],
                total=group[
                    "expected_students"
                ],
            )
        )

    result.sort(
        key=lambda item: (
            item["started_at"] or ""
        ),
        reverse=True,
    )

    return result


# ============================================================
# BUILD COMPLETE ADMIN ATTENDANCE REPORT
# ============================================================

def get_admin_attendance_report(
    db: Session,
    semester_id: int | None = None,
    section_id: int | None = None,
    subject_id: int | None = None,
    faculty_id: int | None = None,
    date_from=None,
    date_to=None,
) -> dict:
    if (
        date_from is not None
        and date_to is not None
        and date_from > date_to
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "date_from cannot be later than date_to."
            ),
        )

    rows = _get_admin_report_rows(
        db=db,
        semester_id=semester_id,
        section_id=section_id,
        subject_id=subject_id,
        faculty_id=faculty_id,
        date_from=date_from,
        date_to=date_to,
    )

    return {
        "filters": {
            "semester_id": semester_id,
            "section_id": section_id,
            "subject_id": subject_id,
            "faculty_id": faculty_id,
            "date_from": (
                date_from.isoformat()
                if date_from is not None
                else None
            ),
            "date_to": (
                date_to.isoformat()
                if date_to is not None
                else None
            ),
        },

        "summary":
            _admin_report_summary(
                rows
            ),

        "faculty_wise":
            _admin_faculty_wise_report(
                rows
            ),

        "subject_wise":
            _admin_subject_wise_report(
                rows
            ),

        "section_wise":
            _admin_section_wise_report(
                rows
            ),

        "day_wise":
            _admin_day_wise_report(
                rows
            ),

        "session_wise":
            _admin_session_wise_report(
                rows
            ),

        "records":
            rows,
    }