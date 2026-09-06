"""
Subject Catalog Export Service
==============================

Build CSV, Excel and PDF exports for Admin Subject Management.

The export follows the same Program -> Branch -> Semester + status +
search filters used by the Subject Management UI.
"""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


EXPORT_HEADERS = [
    "Program Code",
    "Program Name",
    "Branch Code",
    "Branch Name",
    "Semester Number",
    "Subject Code",
    "Subject Name",
    "Credits",
    "Status",
]


def _credit_value(value) -> float:
    if value is None:
        return 0.0

    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _display_credit(value) -> int | float:
    number = _credit_value(value)
    if number.is_integer():
        return int(number)
    return number


def _subject_export_row(subject: dict) -> list:
    return [
        subject.get("program_code") or "",
        subject.get("program_name") or "",
        subject.get("branch_code") or "",
        subject.get("branch_name") or "",
        subject.get("semester_number") or "",
        subject.get("subject_code") or "",
        subject.get("subject_name") or "",
        _display_credit(subject.get("credits")),
        "ACTIVE" if subject.get("is_active") else "INACTIVE",
    ]


def filter_subject_export_rows(
    subjects: Iterable[dict],
    *,
    status_filter: str = "ALL",
    search: str | None = None,
) -> list[dict]:
    normalized_status = (status_filter or "ALL").strip().upper()

    if normalized_status not in {"ALL", "ACTIVE", "INACTIVE"}:
        normalized_status = "ALL"

    query = (search or "").strip().lower()

    result: list[dict] = []

    for subject in subjects:
        active = bool(subject.get("is_active"))

        if normalized_status == "ACTIVE" and not active:
            continue

        if normalized_status == "INACTIVE" and active:
            continue

        if query:
            searchable = " ".join(
                str(subject.get(key) or "")
                for key in (
                    "subject_id",
                    "program_code",
                    "program_name",
                    "branch_code",
                    "branch_name",
                    "semester_number",
                    "semester_name",
                    "subject_code",
                    "subject_name",
                    "credits",
                )
            ).lower()

            if query not in searchable:
                continue

        result.append(subject)

    return result


def _summary(subjects: list[dict]) -> tuple[int, float]:
    return (
        len(subjects),
        sum(_credit_value(item.get("credits")) for item in subjects),
    )


def build_subject_catalog_csv(subjects: list[dict]) -> bytes:
    output = StringIO()
    writer = csv.writer(output)

    writer.writerow(EXPORT_HEADERS)

    for subject in subjects:
        writer.writerow(_subject_export_row(subject))

    # UTF-8 BOM helps Excel open the CSV correctly.
    return output.getvalue().encode("utf-8-sig")


def build_subject_catalog_excel(subjects: list[dict]) -> bytes:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Subjects"

    sheet.append(EXPORT_HEADERS)

    for subject in subjects:
        sheet.append(_subject_export_row(subject))

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = {
        1: 16,  # Program Code
        2: 28,  # Program Name
        3: 16,  # Branch Code
        4: 32,  # Branch Name
        5: 18,  # Semester Number
        6: 20,  # Subject Code
        7: 52,  # Subject Name
        8: 12,  # Credits
        9: 12,  # Status
    }

    for index, width in widths.items():
        sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # --------------------------------------------------------
    # Summary sheet: Program -> Branch -> Semester
    # --------------------------------------------------------
    summary_sheet = workbook.create_sheet("Summary")
    summary_sheet.append([
        "Program",
        "Branch",
        "Semester",
        "Number of Subjects",
        "Total Credits",
    ])

    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    groups: dict[tuple, list[dict]] = {}

    for subject in subjects:
        key = (
            subject.get("program_code") or "",
            subject.get("branch_code") or "",
            subject.get("semester_number") or "",
        )
        groups.setdefault(key, []).append(subject)

    for key in sorted(
        groups,
        key=lambda item: (
            str(item[0]).lower(),
            str(item[1]).lower(),
            int(item[2]) if str(item[2]).isdigit() else 999,
        ),
    ):
        items = groups[key]
        count, credits = _summary(items)
        summary_sheet.append([
            key[0],
            key[1],
            key[2],
            count,
            _display_credit(credits),
        ])

    total_count, total_credits = _summary(subjects)

    summary_sheet.append([])
    summary_sheet.append([
        "TOTAL",
        "",
        "",
        total_count,
        _display_credit(total_credits),
    ])

    summary_sheet.freeze_panes = "A2"

    for row in summary_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    for column, width in {
        1: 18,
        2: 18,
        3: 14,
        4: 20,
        5: 16,
    }.items():
        summary_sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_subject_catalog_pdf(
    subjects: list[dict],
    *,
    scope_label: str = "All Subjects",
) -> bytes:
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title="Subject Catalog",
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(
        Paragraph(
            "<b>Subject Management Export</b>",
            styles["Title"],
        )
    )
    story.append(
        Paragraph(
            scope_label,
            styles["Normal"],
        )
    )

    count, credits = _summary(subjects)

    story.append(
        Paragraph(
            f"{count} subject(s) &nbsp;&nbsp; | &nbsp;&nbsp; "
            f"{_display_credit(credits)} total credits",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 5 * mm))

    pdf_headers = [
        "Program",
        "Branch",
        "Sem",
        "Subject Code",
        "Subject Name",
        "Credits",
        "Status",
    ]

    data = [pdf_headers]

    body_style = styles["BodyText"]
    body_style.fontSize = 7
    body_style.leading = 8
    body_style.alignment = 1

    for subject in subjects:
        data.append([
            Paragraph(
                str(subject.get("program_code") or ""),
                body_style,
            ),
            Paragraph(
                str(subject.get("branch_code") or ""),
                body_style,
            ),
            str(subject.get("semester_number") or ""),
            Paragraph(
                str(subject.get("subject_code") or ""),
                body_style,
            ),
            Paragraph(
                str(subject.get("subject_name") or ""),
                body_style,
            ),
            str(_display_credit(subject.get("credits"))),
            "ACTIVE" if subject.get("is_active") else "INACTIVE",
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            25 * mm,
            25 * mm,
            13 * mm,
            34 * mm,
            105 * mm,
            18 * mm,
            22 * mm,
        ],
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1F4E78"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER",
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.25,
                colors.HexColor("#B7C9D6"),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F7FAFC"),
                ],
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
        ])
    )

    story.append(table)

    document.build(story)

    return output.getvalue()
